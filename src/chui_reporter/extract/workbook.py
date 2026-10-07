"""Anchor-based reader for human-formatted finance workbooks.

The source workbooks are not tabular data; they are accountants' working papers.
Row and column positions drift between quarters and between the GP and LP books,
so nothing here addresses cells positionally at the call site. Callers name a
label ("Total Net Asset Value", "5100") and read a fixed offset from wherever
that label turns out to be.

Three rules this module enforces, all learned from the Q1/Q2 source files:

1. Cached values only. Every workbook ships a calcChain, so the cached value is
   what the fund administrator actually saw. We never re-evaluate formulas --
   several sheets contain TODAY(), which would drift on every read.
2. Excel errors are quarantined, never coerced. A #REF! that silently becomes 0
   is how a wrong number reaches an LP report.
3. Ambiguity raises. An anchor that matches zero times or more than once is a
   bug in the extraction rule, not something to guess past.
"""

from __future__ import annotations

import gc
import hashlib
import os
import re
import threading
import unicodedata
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

EXCEL_ERRORS = frozenset(
    {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#GETTING_DATA"}
)

# Excel's 1900 date system, with the epoch shifted to absorb the fictional 1900 leap day.
EXCEL_EPOCH = date(1899, 12, 30)

Unit = Literal[
    "USD", "USD_thousands", "USD_millions", "ratio", "percent", "count", "months", "text"
]


class ExtractionError(Exception):
    """Raised when an extraction rule cannot be satisfied unambiguously."""


class AnchorNotFound(ExtractionError):
    pass


class AnchorAmbiguous(ExtractionError):
    pass


@dataclass(frozen=True)
class Provenance:
    """Where a value came from. Every fact carries one; none are optional."""

    file_path: str
    file_sha256: str
    sheet: str
    cell: str
    anchor_label: str | None = None
    anchor_cell: str | None = None

    def __str__(self) -> str:
        where = f"{Path(self.file_path).name}::{self.sheet}!{self.cell}"
        if self.anchor_label:
            where += f" (anchor {self.anchor_label!r} @ {self.anchor_cell})"
        return where


@dataclass(frozen=True)
class Value:
    """An extracted value and everything needed to defend it."""

    raw: Any
    unit: Unit
    provenance: Provenance
    status: Literal["extracted", "quarantined"] = "extracted"
    note: str | None = None

    @property
    def is_usable(self) -> bool:
        return self.status == "extracted" and self.raw is not None

    def as_usd(self) -> float:
        """Normalise to absolute USD. Unit is declared by the extraction rule,
        never inferred -- several sheets put thousands in a column labelled USD."""
        if not self.is_usable:
            raise ExtractionError(f"value not usable ({self.status}): {self.provenance}")
        if not isinstance(self.raw, (int, float)):
            raise ExtractionError(f"not numeric: {self.raw!r} at {self.provenance}")
        factor = {"USD": 1, "USD_thousands": 1_000, "USD_millions": 1_000_000}.get(self.unit)
        if factor is None:
            raise ExtractionError(f"unit {self.unit!r} is not a currency at {self.provenance}")
        return float(self.raw) * factor


def normalise_label(text: Any) -> str:
    """Canonical form for label matching.

    The corpus contains trailing spaces, non-breaking spaces, doubled spaces and
    smart punctuation in otherwise identical labels, so comparisons are made on a
    folded form rather than the raw string.
    """
    if text is None:
        return ""
    s = unicodedata.normalize("NFKC", str(text))
    s = s.replace(" ", " ").replace("’", "'").replace("‘", "'")
    s = s.replace("“", '"').replace("”", '"')
    s = s.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", s).strip().casefold()


def excel_serial_to_date(serial: float) -> date:
    """Excel serial -> date. Period headers are first-of-month serials formatted
    mmm-yy; 46174 renders 'Jun-26' and means the period ended 30 June 2026."""
    return EXCEL_EPOCH + timedelta(days=int(serial))


@dataclass
class Sheet:
    """A worksheet with merged ranges flattened and labels indexed."""

    name: str
    _ws: Worksheet
    _book: Workbook
    _grid: dict[tuple[int, int], Any] = field(default_factory=dict, repr=False)
    _labels: dict[str, list[tuple[int, int]]] = field(default_factory=dict, repr=False)

    def __post_init__(self) -> None:
        self._build_grid()
        self._index_labels()

    def _build_grid(self) -> None:
        """Flatten the sheet, propagating merged values across their whole range.

        A merged label in these workbooks covers several cells but openpyxl
        exposes the value only on the top-left; every other covered cell reads
        None. Propagating is the single highest-value normalisation step for
        human-formatted sheets.
        """
        for row in self._ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    self._grid[(cell.row, cell.column)] = cell.value
        for rng in self._ws.merged_cells.ranges:
            anchor = self._grid.get((rng.min_row, rng.min_col))
            if anchor is None:
                continue
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self._grid.setdefault((r, c), anchor)

    def _index_labels(self) -> None:
        for (r, c), v in self._grid.items():
            if isinstance(v, str):
                key = normalise_label(v)
                if key:
                    self._labels.setdefault(key, []).append((r, c))

    @staticmethod
    def addr(row: int, col: int) -> str:
        return f"{get_column_letter(col)}{row}"

    def find(self, label: str, *, column: int | None = None) -> tuple[int, int]:
        """Locate a label exactly once. Raises on zero or multiple matches.

        `column` restricts the search to one column, which is how account-code
        lookups stay unambiguous when the same text appears in a note elsewhere.
        """
        hits = self._labels.get(normalise_label(label), [])
        if column is not None:
            hits = [h for h in hits if h[1] == column]
        if not hits:
            raise AnchorNotFound(f"{label!r} not found in {self._book.path.name}::{self.name}")
        if len(hits) > 1:
            locs = ", ".join(self.addr(r, c) for r, c in sorted(hits))
            raise AnchorAmbiguous(
                f"{label!r} matched {len(hits)} cells in "
                f"{self._book.path.name}::{self.name} ({locs}) -- narrow the rule"
            )
        return hits[0]

    def find_prefix(self, prefix: str, *, column: int | None = None) -> list[tuple[int, int]]:
        """All labels starting with `prefix`. Used for account codes written as
        '5100 (Partnership Interest)', where the code is stable and the name is not."""
        want = normalise_label(prefix)
        out = [
            (r, c)
            for key, locs in self._labels.items()
            if key.startswith(want)
            for (r, c) in locs
            if column is None or c == column
        ]
        return sorted(out)

    def raw(self, row: int, col: int) -> Any:
        return self._grid.get((row, col))

    def value(
        self,
        anchor: str,
        *,
        dx: int = 0,
        dy: int = 0,
        unit: Unit = "USD",
        column: int | None = None,
    ) -> Value:
        """Read the cell at a fixed offset from `anchor`.

        Offsets are relative because absolute addresses drift; the label does not.
        """
        ar, ac = self.find(anchor, column=column)
        return self.at(ar + dy, ac + dx, unit=unit, anchor_label=anchor, anchor_cell=self.addr(ar, ac))

    def at(
        self,
        row: int,
        col: int,
        *,
        unit: Unit = "USD",
        anchor_label: str | None = None,
        anchor_cell: str | None = None,
    ) -> Value:
        raw = self.raw(row, col)
        prov = Provenance(
            file_path=str(self._book.path),
            file_sha256=self._book.sha256,
            sheet=self.name,
            cell=self.addr(row, col),
            anchor_label=anchor_label,
            anchor_cell=anchor_cell,
        )
        if isinstance(raw, str) and raw.strip() in EXCEL_ERRORS:
            return Value(
                raw=None,
                unit=unit,
                provenance=prov,
                status="quarantined",
                note=f"Excel error {raw.strip()} -- not coerced",
            )
        if isinstance(raw, datetime):
            raw = raw.date()
        return Value(raw=raw, unit=unit, provenance=prov)

    def column_of(self, header: str) -> int:
        return self.find(header)[1]

    def __repr__(self) -> str:  # pragma: no cover - debugging aid
        return f"<Sheet {self.name!r} cells={len(self._grid)}>"


# ---- memory ---------------------------------------------------------------------------------------------------------------
# A workbook held in memory costs about a hundred times its size on disk, and the server this runs on can be small (512 MB, of
# which the application itself uses over 100). So workbooks are opened through `Workbook.open`, which keeps a few of them, shared,
# and lets go of the least recently used before it would run the server out of memory, and refuses to open one when what is
# left would not hold it (an error the agent can read) rather than letting the whole server crawl.
_MB = 1_000_000
_PER_BYTE = 100
_BUDGET = int(float(os.environ.get("CHUI_WORKBOOK_CACHE_MB", "220")) * _MB)       # estimated memory the open workbooks may hold
_LOAD = threading.Lock()
_OPEN: OrderedDict = OrderedDict()


def headroom() -> int | None:
    """Bytes this container can still use before its memory limit (cache that can be dropped does not count), or None when
    the container does not say."""
    try:
        base = Path("/sys/fs/cgroup")
        if (base / "memory.max").exists():
            raw = (base / "memory.max").read_text().strip()
            if raw == "max":
                return None
            limit, used = int(raw), int((base / "memory.current").read_text())
            stat = dict(ln.split()[:2] for ln in (base / "memory.stat").read_text().splitlines() if ln.strip())
            used -= int(stat.get("inactive_file", 0))
        else:
            v1 = base / "memory"
            limit, used = int((v1 / "memory.limit_in_bytes").read_text()), int((v1 / "memory.usage_in_bytes").read_text())
            stat = dict(ln.split()[:2] for ln in (v1 / "memory.stat").read_text().splitlines() if ln.strip())
            used -= int(stat.get("total_inactive_file", 0))
        return None if limit > 1 << 50 else limit - used
    except (OSError, ValueError, IndexError):
        return None


def _held() -> int:
    return sum(w._est for w in _OPEN.values())


def _evict(need: int) -> None:
    """Let go of the least recently used workbooks until `need` more would fit in the budget."""
    dropped = False
    while _OPEN and _held() + need > _BUDGET:
        _OPEN.popitem(last=False)
        dropped = True
    if dropped:
        gc.collect()                            # the sheets point back at their book, so reference counting alone does not free it


def release_all() -> None:
    """Let go of every shared workbook (a heavy job is over)."""
    with _LOAD:
        _evict(_BUDGET + 1)


class Workbook:
    """A source workbook, opened for cached values only."""

    @classmethod
    def open(cls, path: str | Path) -> "Workbook":
        """The shared, memory-bounded way to open a workbook: the same object for the same unchanged file, a few at a time."""
        p = Path(path)
        st = p.stat()                           # FileNotFoundError, as the constructor would
        key = (str(p.resolve()), st.st_size, st.st_mtime_ns)
        need = st.st_size * _PER_BYTE
        with _LOAD:
            wb = _OPEN.get(key)
            if wb is not None:
                _OPEN.move_to_end(key)
                return wb
            _evict(need)
            room = headroom()
            if room is not None and room < need * 1.3:
                _evict(_BUDGET + 1)
                room = headroom()
                if room is not None and room < need * 1.3:
                    raise ExtractionError(
                        f"There is not enough memory on this server to open {p.name} right now (about {room // _MB} MB free, "
                        f"about {int(need * 1.3) // _MB} MB needed). Other work is using it: wait for that to finish, then try again.")
            wb = cls(p)
            wb._shared, wb._est = True, need
            _OPEN[key] = wb
            return wb

    _shared = False
    _est = 0

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(self.path)
        self.sha256 = _sha256(self.path)
        # data_only=True yields the cached value; we never evaluate formulas.
        self._wb = load_workbook(self.path, data_only=True, read_only=False)
        # Sheet names in this corpus carry trailing spaces ("Fund expense ", "IRR ").
        self._sheets: dict[str, str] = {normalise_label(n): n for n in self._wb.sheetnames}
        self._cache: dict[str, Sheet] = {}

    @property
    def sheet_names(self) -> list[str]:
        return list(self._wb.sheetnames)

    def sheet(self, name: str) -> Sheet:
        key = normalise_label(name)
        if key not in self._sheets:
            raise ExtractionError(
                f"sheet {name!r} not in {self.path.name}; available: {self.sheet_names}"
            )
        real = self._sheets[key]
        if real not in self._cache:
            self._cache[real] = Sheet(name=real, _ws=self._wb[real], _book=self)
        return self._cache[real]

    def has_sheet(self, name: str) -> bool:
        return normalise_label(name) in self._sheets

    def close(self) -> None:
        if not self._shared:                    # a shared workbook is other callers' too; it is let go of when evicted
            self._wb.close()

    def __enter__(self) -> Workbook:
        return self

    def __exit__(self, *exc: object) -> None:
        self.close()

    def __repr__(self) -> str:  # pragma: no cover - debugging aid
        return f"<Workbook {self.path.name!r} sheets={len(self._sheets)}>"


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()
