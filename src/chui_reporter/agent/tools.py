"""Tools exposed to the ReAct agent.

Each tool returns text the model can reason over, and every numeric tool returns
provenance alongside the figure so the agent can cite a source rather than
recall one. The extraction layer underneath is the same code the tests cover.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Annotated

from langchain_core.tools import InjectedToolCallId, tool
from langgraph.types import interrupt

from .. import config
from .. import period as pr
from ..extract.published_report import read_portfolio_performance
from ..extract.valuation import Quarter, ValuationWorkbook, discover
from ..extract.workbook import Workbook
from ..extract.workpaper import StatementReader, capital_calls, derive_fund_metrics
from ..validate.reconcile import reconcile
from .ledger import DerivationError, derive_fact, derived, fact_from_value, verify_claim
from . import builders as B
from .store import Fact, Store


_STORE: Store | None = None


def get_store() -> Store:
    """Created on first use, so importing the tools never opens a connection."""
    global _STORE
    if _STORE is None:
        _STORE = Store()
    return _STORE


def set_store(store: Store | None) -> None:
    """Point the tools at a specific store (the web runtime has one report per period)."""
    global _STORE
    _STORE = store


# The agent may ask the user, but not endlessly: a few well-chosen questions per run.
QUESTION_BUDGET = 3
_asked: set[str] = set()


def reset_questions() -> None:
    _asked.clear()

# --------------------------------------------------------------------------
# Source discovery and Excel reading
# --------------------------------------------------------------------------


@tool
def list_sources() -> str:
    """What the agent has to work with: each kind of source document with whether it is
    available, which report sections it feeds, and the files found.

    Use this first, so you know what exists and what is missing before you start."""
    from ..workspace.slots import SLOTS, slot_files

    paths = [p.relative_to(config.SOURCE_ROOT).as_posix() for p in config.SOURCE_ROOT.rglob("*")
             if p.is_file() and not p.name.startswith((".", "~$"))]
    out = []
    for slot in SLOTS:
        files = slot_files(slot, paths)
        state = "AVAILABLE" if len(files) >= slot.minimum else ("MISSING (required)" if slot.required
                                                                  else "not provided (optional)")
        out.append(f"{slot.label} -- {state}; feeds sections {', '.join(slot.unlocks)}")
        for f in files[:12]:
            out.append(f"    {f}")
        if len(files) > 12:
            out.append(f"    ... and {len(files) - 12} more")
    attached = sorted(p for p in paths if p.startswith("Uploads/"))
    if attached:
        out.append("Attached by the user in the conversation (read them if the user's message refers to them):")
        out.extend(f"    {f}" for f in attached[:20])
    return "\n".join(out)


@tool
def look_at_image(file_name: str, question: str = "Describe what this shows.") -> str:
    """Look at an image the user attached (PNG, JPG, WEBP, GIF) and answer a question about it.

    Use it for layout, design and qualitative guidance (a screenshot of a problem, a sketch of what
    they want, a photograph). Figures that appear only in an image CANNOT be recorded in the fact
    ledger or used in the report, because nothing can verify them: if the user wants numbers from an
    image, ask for the source document instead."""
    import base64
    import io

    from langchain_core.messages import HumanMessage

    from PIL import Image

    from .llm import get_llm

    try:
        path = _resolve_typed(file_name, (".png", ".jpg", ".jpeg", ".webp", ".gif"))
    except FileNotFoundError as exc:
        return f"ERROR: {exc}"
    if path.suffix.lower() not in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
        return f"ERROR: {path.name} is not an image; use read_pdf, read_text or the excel tools."
    try:
        img = Image.open(path).convert("RGB")
        img.thumbnail((1600, 1600))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=85)
        url = "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
        reply = get_llm().invoke([HumanMessage(content=[
            {"type": "text", "text": question[:600]}, {"type": "image_url", "image_url": {"url": url}}])])
    except Exception as exc:  # noqa: BLE001
        return f"ERROR: could not look at {path.name}: {type(exc).__name__}: {exc}"
    return (f"[{path.name}] {reply.content}\n(Figures read from an image are not verifiable and must not "
            f"be recorded or used in the report.)")


@tool
def read_text(file_name: str, max_chars: int = 6000) -> str:
    """Read a plain-text, CSV, Markdown or Word (.docx) source document: macro indicator
    files, a GP statement, notes. For PDFs use read_pdf; for workbooks use the excel tools.

    Returns the text (tables in a .docx are flattened row by row). Long files are cut at
    `max_chars`; ask for more by raising it (at most 20000)."""
    try:
        path = _resolve_typed(file_name, (".txt", ".csv", ".md", ".docx", ".json"))
    except FileNotFoundError as exc:
        return f"ERROR: {exc}"
    suffix = path.suffix.lower()
    if suffix in (".pdf",):
        return f"ERROR: {path.name} is a PDF; use read_pdf."
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return f"ERROR: {path.name} is a workbook; use excel_sheets / excel_dump_region."
    try:
        if suffix == ".docx":
            from docx import Document

            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for t in doc.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        return f"ERROR: could not read {path.name}: {type(exc).__name__}: {exc}"
    limit = max(500, min(int(max_chars), 20000))
    shown = text[:limit]
    more = f"\n[... {len(text) - limit:,} more characters; raise max_chars to read on]" if len(text) > limit else ""
    return f"[{path.name}, {len(text):,} characters]\n{shown}{more}"


@tool
def read_pdf(file_name: str, first_page: int = 1, last_page: int = 3) -> str:
    """Read the text of a PDF source document, with table columns kept aligned.

    Use this for anything that is a PDF: the Uncover drawdown request in
    'Pipeline and subsequent events', or a prior-period report in 'Prior Period
    Baseline'. `file_name` may be a partial name. Reads at most 6 pages per call.
    """
    import re as _re

    from pypdf import PdfReader

    path = _resolve_typed(file_name, (".pdf",))
    if path.suffix.lower() != ".pdf":
        return f"ERROR: {path.name} is not a PDF; use the excel_* tools for workbooks."
    reader = PdfReader(str(path))
    total = len(reader.pages)
    first = max(1, first_page)
    last = min(total, last_page, first + 5)
    out = [f"{path.name} -- {total} pages; showing {first}-{last}"]
    for n in range(first, last + 1):
        text = reader.pages[n - 1].extract_text(extraction_mode="layout") or ""
        text = _re.sub(r"[ \t]{3,}", "  ", text)
        text = "\n".join(ln.rstrip() for ln in text.splitlines() if ln.strip())
        out.append(f"--- page {n} ---\n{text[:6000]}")
    return "\n".join(out)


def _require_excel(path: Path) -> str | None:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return (f"ERROR: {path.name} is not an Excel workbook"
                + (" -- use read_pdf for PDFs." if path.suffix.lower() == ".pdf" else "."))
    return None


@tool
def excel_sheets(file_name: str) -> str:
    """List the sheet names in a source workbook. `file_name` may be a partial name."""
    path = _resolve_typed(file_name, (".xlsx", ".xlsm"))
    if (err := _require_excel(path)):
        return err
    wb = Workbook(path)
    names = wb.sheet_names
    wb.close()
    return f"{path.name}: {names}"


@tool
def excel_find_value(file_name: str, sheet: str, label: str,
                     dx: int = 1, dy: int = 0) -> str:
    """Find a labelled cell in a workbook and read a value at an offset from it.

    Values are located by label, not cell address, because rows shift between
    quarters. Returns the value with its sheet, cell and anchor so it can be
    cited. `dx` moves right, `dy` moves down from the label.
    """
    path = _resolve_typed(file_name, (".xlsx", ".xlsm"))
    if (err := _require_excel(path)):
        return err
    wb = Workbook(path)
    try:
        v = wb.sheet(sheet).value(label, dx=dx, dy=dy, unit="USD")
        return (f"value={v.raw!r} status={v.status} "
                f"source={v.provenance}" + (f" note={v.note}" if v.note else ""))
    except Exception as exc:  # noqa: BLE001 - the agent should see the failure
        return f"ERROR: {exc}"
    finally:
        wb.close()


@tool
def excel_dump_region(file_name: str, sheet: str, first_row: int = 1,
                      last_row: int = 40, first_col: int = 1, last_col: int = 10) -> str:
    """Dump a rectangular region of a sheet as text, to inspect its layout."""
    path = _resolve_typed(file_name, (".xlsx", ".xlsm"))
    if (err := _require_excel(path)):
        return err
    wb = Workbook(path)
    try:
        sh = wb.sheet(sheet)
        lines = []
        for r in range(first_row, last_row + 1):
            cells = [f"{sh.addr(r, c)}={sh.raw(r, c)!r}"
                     for c in range(first_col, last_col + 1) if sh.raw(r, c) is not None]
            if cells:
                lines.append(" | ".join(cells))
        return "\n".join(lines) or "(region empty)"
    finally:
        wb.close()


# --------------------------------------------------------------------------
# Domain tools -- the reconciled views
# --------------------------------------------------------------------------


def _fact_at(label: str, value: float, prov, *, note: str | None = None,
             unit: str = "USD") -> Fact:
    """A ledger fact whose provenance was supplied by an extractor, not the model."""
    return Fact(label, float(value), unit=unit, source_file=Path(prov.file_path).name,
                source_sheet=prov.sheet, source_cell=prov.cell, status="extracted", note=note)


def _recorded(facts: list[Fact]) -> str:
    n = get_store().add_facts([f for f in facts if f is not None])
    return f"[ledger: {n} facts recorded by this tool with source provenance; do not re-save them]"


@tool
def portfolio_valuations(year: int = 0, quarter: int = 0) -> str:
    """Reconciled per-company fair values for a quarter.

    Compares the per-company valuation workbooks against the Fund Model and
    flags disagreements, carried-forward (stale) marks, and unit mislabelling.
    These two sources genuinely disagree, so prefer this over reading either
    source alone. Every figure shown is also recorded in the fact ledger, with
    the source cell, under the labels listed at the end -- including each
    company's difference and the portfolio totals, computed in code.
    """
    P = pr.current()
    target = Quarter(year or P.year, quarter or P.q)
    rows, dis = reconcile(config.VALUATION_REPORTS, config.fund_model(), target)
    out = ["company | valuation_report_usd | fund_model_usd | stale_quarters | flags"]
    facts: list[Fact] = []
    tot_vr = tot_fm = 0.0
    for r in rows:
        out.append(
            f"{r.company} | {_fmt(r.vr_fair_value_usd)} | {_fmt(r.fm_fair_value_usd)} | "
            f"{r.staleness} | {'; '.join(r.flags) or '-'}"
            + (f" | UNIT: {r.unit_conflict}" if r.unit_conflict else "")
        )
        lab_vr = f"{r.company} fair value (valuation report)"
        lab_fm = f"{r.company} fair value (fund model)"
        if r.vr_fair_value_usd is not None and r.vr_prov is not None:
            facts.append(_fact_at(lab_vr, r.vr_fair_value_usd, r.vr_prov,
                                  note=f"mark from {r.vr_quarter}; staleness {r.staleness}q"))
            tot_vr += r.vr_fair_value_usd
        if r.fm_fair_value_usd is not None and r.fm_prov is not None:
            facts.append(_fact_at(lab_fm, r.fm_fair_value_usd, r.fm_prov,
                                  note="Fund Model 'Portfolio Valuation', US$000s x 1000"))
            tot_fm += r.fm_fair_value_usd
        if r.vr_fair_value_usd is not None and r.fm_fair_value_usd is not None:
            facts.append(derived(f"{r.company} fair value difference (valuation report minus fund model)",
                                 r.vr_fair_value_usd - r.fm_fair_value_usd, [lab_vr, lab_fm],
                                 "valuation report - fund model"))
    out.append("")
    out.append(f"{len(dis)} material disagreements:")
    for d in dis:
        out.append(f"  {d.company}: valuation_report={_fmt(d.valuation_report)} "
                   f"fund_model={_fmt(d.fund_model)} {d.note}")
    facts += [
        derived("Portfolio fair value total (valuation reports)", tot_vr,
                [f"{r.company} fair value (valuation report)" for r in rows
                 if r.vr_fair_value_usd is not None], "sum"),
        derived("Portfolio fair value total (fund model)", tot_fm,
                [f"{r.company} fair value (fund model)" for r in rows
                 if r.fm_fair_value_usd is not None], "sum"),
        derived("Portfolio fair value total difference (valuation reports minus fund model)",
                tot_vr - tot_fm, ["Portfolio fair value total (valuation reports)",
                                  "Portfolio fair value total (fund model)"], "difference"),
    ]
    out.append(f"\nTotals: valuation reports {tot_vr:,.2f} | fund model {tot_fm:,.2f} "
               f"| difference {tot_vr - tot_fm:,.2f}")
    out.append("Ledger labels: '<Company> fair value (valuation report)', "
               "'<Company> fair value (fund model)', "
               "'<Company> fair value difference (valuation report minus fund model)', "
               "'Portfolio fair value total (valuation reports|fund model)', "
               "'Portfolio fair value total difference (valuation reports minus fund model)'")
    out.append(_recorded(facts))
    return "\n".join(out)


_CALL_PROV = re.compile(r"^(?P<file>.*?)::(?P<sheet>.*?)!(?P<cell>[A-Z]+\d+)")


@tool
def fund_capital_position() -> str:
    """Capital calls and derived fund metrics for the Mauritius LP vehicle.

    Note the workpapers contain NO fund-level performance metrics -- no NAV,
    TVPI, DPI or IRR line exists and investments are carried at cost. Everything
    here is derived, and each figure says how. All figures are recorded in the
    fact ledger (labels listed at the end).
    """
    wb = Workbook(config.lp_workpaper())
    try:
        calls = capital_calls(wb)
        basis = wb.sheet("Management fee").value("capital committed by MEDA", dx=1, unit="USD")
        m = derive_fund_metrics(wb, commitment_usd=basis.as_usd() if basis.is_usable else None)
        facts: list[Fact] = []
        lines = ["Capital calls (LP Mauritius, counterparty MEDA Mauritius Foundation):"]
        for c in calls:
            lines.append(f"  #{c.sequence} {str(c.date)[:10]} ${c.amount_usd:,.2f}")
            mm = _CALL_PROV.match(c.provenance)
            facts.append(Fact(f"Capital call {c.sequence} amount (LP Mauritius)", c.amount_usd,
                              unit="USD", source_file=Path(mm["file"]).name if mm else None,
                              source_sheet=mm["sheet"] if mm else None,
                              source_cell=mm["cell"] if mm else None,
                              as_of=str(c.date)[:10], status="extracted"))
        total = sum(c.amount_usd for c in calls)
        lines.append(f"  TOTAL ${total:,.2f}")
        facts.append(derived("Total capital called (LP Mauritius)", total,
                             [f"Capital call {c.sequence} amount (LP Mauritius)" for c in calls], "sum"))
        if basis.is_usable:
            facts.append(_fact_at("Management-fee commitment basis (LP Mauritius)", basis.as_usd(),
                                  basis.provenance,
                                  note="'capital committed by MEDA' -- the fee basis, which is "
                                       "not necessarily the LP commitment of record"))
        lines += ["", "Derived metrics (scope: LP_MAURITIUS only -- one of three vehicles):"]
        metric_labels = {
            "contributed_usd": "Contributed capital (LP Mauritius)",
            "uncalled_usd": "Uncalled capital against the fee basis (LP Mauritius)",
            "nav_cost_basis_usd": "NAV at cost (LP Mauritius)",
            "investments_at_cost_usd": "Investments at cost (LP Mauritius)",
            "cash_usd": "Cash (LP Mauritius)",
            "distributions_usd": "Distributions (LP Mauritius)",
        }
        for k, lab in metric_labels.items():
            val = getattr(m, k)
            if val is None:
                continue
            lines.append(f"  {lab} = {_fmt(val)}")
            facts.append(derived(lab, val, [m.source.get(k, "")], m.source.get(k, "")))
        for lab, val, how in [
            ("Percent of fee basis called (LP Mauritius)", m.pct_called, "contributed / fee basis"),
            ("DPI (LP Mauritius)", m.dpi, "distributions / contributed"),
            ("RVPI cost basis (LP Mauritius)", m.rvpi_cost_basis, "NAV at cost / contributed"),
        ]:
            if val is not None:
                lines.append(f"  {lab} = {val:.4f}")
                facts.append(derived(lab, val, [], how, unit="ratio"))
        lines += ["", "Provenance:"] + [f"  {k}: {v}" for k, v in m.source.items()]
        lines.append(_recorded(facts))
        return "\n".join(lines)
    finally:
        wb.close()


_SOFP_TOTALS = ["Total Investments", "Total Non Current Assets", "Total Other Current Assets",
                "Total Cash and Cash Equivalents", "Total Current Assets", "Total Assets",
                "Total Equity", "Total Other Current Liabilities", "Total Current Liabilities",
                "Total Equity & Liabilities"]
_SOCI_TOTALS = ["Total Other Expense", "Loss Before Tax"]


@tool
def financial_statements() -> str:
    """LP balance sheet and income statement line items, keyed by account code.

    Every line and total, current and comparative, is recorded in the fact
    ledger with its source cell. Labels: 'LP SOFP <code> <name> [current|comparative]',
    'LP SOFP <total label> [current|comparative]', 'LP SOCI <code> <name>',
    'LP SOCI <total label>'.
    """
    from ..extract.workbook import ExtractionError

    wb = Workbook(config.lp_workpaper())
    try:
        r = StatementReader(wb, "LP")
        facts: list[Fact] = []
        lines = ["BALANCE SHEET (SOFP) -- code | label | current_usd | comparative_usd"]
        for code, ln in sorted(r.balance_sheet().items()):
            cur = ln.current.raw if ln.current.is_usable else None
            cmp_ = ln.comparative.raw if ln.comparative and ln.comparative.is_usable else None
            lines.append(f"  {code} | {ln.label} | {_fmt(cur)} | {_fmt(cmp_)}")
            facts += [fact_from_value(f"LP SOFP {code} {ln.label} [current]", ln.current,
                                      as_of=pr.current().end_iso)]
            if ln.comparative is not None and ln.comparative.raw is not None:
                facts += [fact_from_value(f"LP SOFP {code} {ln.label} [comparative]",
                                          ln.comparative, as_of=f"{pr.current().year - 1}-12-31")]
        lines += ["", "BALANCE SHEET TOTALS -- label | current | comparative"]
        for label in _SOFP_TOTALS:
            try:
                cur = r.total("SOFP", label, 3)
                cmp_ = r.total("SOFP", label, 5)
            except ExtractionError:
                continue
            lines.append(f"  {label} | {_fmt(cur.raw)} | {_fmt(cmp_.raw)}")
            facts += [fact_from_value(f"LP SOFP {label} [current]", cur, as_of=pr.current().end_iso),
                      fact_from_value(f"LP SOFP {label} [comparative]", cmp_, as_of=f"{pr.current().year - 1}-12-31")]
        lines += ["", f"INCOME STATEMENT (SOCI), year to {pr.current().end_label} -- code | label | current_usd"]
        for code, ln in sorted(r.income_statement().items()):
            cur = ln.current.raw if ln.current.is_usable else None
            lines.append(f"  {code} | {ln.label} | {_fmt(cur)}")
            facts.append(fact_from_value(f"LP SOCI {code} {ln.label}", ln.current))
        for label in _SOCI_TOTALS:
            try:
                v = r.total("SOCI", label, 3)
            except ExtractionError:
                continue
            lines.append(f"  {label} | {_fmt(v.raw)}")
            facts.append(fact_from_value(f"LP SOCI {label}", v))
        lines.append(_recorded(facts))
        return "\n".join(lines)
    finally:
        wb.close()


@tool
def prior_report_table(year: int = 0, quarter: int = 0) -> str:
    """The Portfolio Performance Summary exactly as published in a prior report.

    Use to obtain PRIOR-quarter comparative columns, and to check what LPs were
    previously told against what the sources now support. Figures from the report
    now being rebuilt are shown for comparison only and are NOT
    recorded in the ledger, so they cannot license a figure.
    """
    P = pr.current()
    year, quarter = year or P.prev.year, quarter or P.prev.q
    pdf = config.prior_report(pr.Period(year, quarter))
    if not pdf.exists():
        return f"ERROR: no published report for Q{quarter} {year} in the workspace"
    rows = read_portfolio_performance(pdf)
    out = ["company | cost_usd | fair_value_100_usd | fair_value_chui_usd | multiple"]
    for p in rows:
        out.append(f"{p.company} | {_fmt(p.cost_usd)} | {_fmt(p.fair_value_100_usd)} | "
                   f"{_fmt(p.fair_value_chui_usd)} | {p.multiple}")
    total = sum(p.fair_value_chui_usd for p in rows)
    out.append(f"TOTAL fair_value_chui_usd = {_fmt(total)}")
    if (year, quarter) == (P.year, P.q):
        out.append("[ledger: NOT recorded -- this is the report being rebuilt; comparison only]")
        return "\n".join(out)
    facts = [Fact(f"{p.company} fair value Chui, as published Q{quarter} {year}",
                  p.fair_value_chui_usd, unit="USD", source_file=pdf.name,
                  source_sheet=f"page {p.page}", status="extracted",
                  note="published prior-period figure") for p in rows]
    facts.append(Fact(f"Portfolio fair value Chui total, as published Q{quarter} {year}", total,
                      unit="USD", source_file=pdf.name, source_sheet=f"page {rows[0].page}",
                      status="extracted", note="sum of published rows"))
    out.append(_recorded(facts))
    return "\n".join(out)


# --------------------------------------------------------------------------
# The mutable report
# --------------------------------------------------------------------------


@tool
def report_save_facts(facts_json: str) -> str:
    """Record a figure you READ from a source with excel_dump_region/excel_find_value/read_pdf.

    Not needed for anything portfolio_valuations, fund_capital_position or
    financial_statements returned -- those record themselves. Each fact needs
    label, value (a plain number in the units you will state it in), unit,
    source_file, and the exact source_sheet + source_cell (workbook) or the page
    number in source_sheet (PDF).

    Each fact is VERIFIED by re-reading the cited cell or page. A fact that does
    not check out, or cannot be checked, is stored as 'claimed' and does NOT
    license any figure in the report. Computed figures belong in
    report_derive_fact, never here.
    """
    try:
        payload = json.loads(facts_json)
    except json.JSONDecodeError as exc:
        return f"ERROR: invalid JSON: {exc}"
    if not isinstance(payload, list):
        return "ERROR: facts_json must be a JSON list of fact objects."
    facts, problems = [], []
    for i, f in enumerate(payload):
        if not isinstance(f, dict) or not f.get("label"):
            problems.append(f"fact #{i}: needs an object with a 'label'")
            continue
        try:
            value = _as_number(f.get("value"))
        except ValueError:
            problems.append(
                f"fact {f['label']!r}: value {f.get('value')!r} is not a number. "
                f"Put dates and text in 'text_value'; 'value' is for figures only."
            )
            continue
        clean = {k: v for k, v in f.items() if k in Fact.__dataclass_fields__}
        clean["value"] = value
        clean.pop("status", None)
        facts.append(Fact(**clean))
    if problems:
        return "ERROR: nothing saved.\n  " + "\n  ".join(problems)

    verified, claimed = [], []
    for f in facts:
        ok, why = verify_claim(f, _resolve) if f.value is not None else (False, "no numeric value")
        f.status = "extracted" if ok else "claimed"
        f.note = (f.note + " | " if f.note else "") + ("verified: " if ok else "UNVERIFIED: ") + why
        (verified if ok else claimed).append((f, why))
    get_store().add_facts([f for f, _ in verified + claimed])
    msg = (f"{len(verified)} verified and recorded; {len(claimed)} stored as CLAIMED "
           f"(they do not license any figure).")
    for f, why in claimed:
        msg += f"\n  not verified: {f.label!r}: {why}"
    return msg


@tool
def report_derive_fact(label: str, expression: str, inputs_json: str, unit: str = "USD") -> str:
    """Compute a figure from facts already in the ledger -- never type a computed number.

    `inputs_json` maps short names to ledger labels, e.g.
    {"a": "Lami fair value (valuation report)", "b": "Lami fair value (fund model)"}
    and `expression` is arithmetic over them (+ - * / and parentheses), e.g. "a - b".
    Every input must be a grounded fact. The result is recorded as 'derived'.
    """
    if len(label) > 90 or any(ch in label for ch in '"{}\n'):
        return ("ERROR: label must be a short plain name (no quotes, braces or newlines, "
                "max 90 characters). It looks like arguments ran into the label.")
    try:
        inputs = json.loads(inputs_json)
        assert isinstance(inputs, dict)
    except Exception:  # noqa: BLE001
        return "ERROR: inputs_json must be a JSON object of name -> ledger label."
    try:
        fact = derive_fact(get_store(), label, expression, inputs, unit)
    except DerivationError as exc:
        return f"ERROR: {exc}"
    get_store().add_facts([fact])
    return f"derived and recorded: {label} = {fact.value:,.6g}   ({expression})"


def _as_number(v) -> float | None:
    """A figure, or None. Dates and prose are rejected rather than coerced --
    the ledger's numeric column is what licenses a number in the report."""
    import math

    if v is None or v == "":
        return None
    if isinstance(v, bool):
        raise ValueError("bool")
    if isinstance(v, (int, float)):
        out = float(v)
    else:
        out = float(str(v).replace(",", "").strip())  # raises ValueError on dates/text
    if not math.isfinite(out):
        raise ValueError("non-finite")
    return out


# --------------------------------------------------------------------------
# Deterministic builders: finished, fully grounded tables and charts
# --------------------------------------------------------------------------


def _build(fn, *a) -> str:
    try:
        return fn(get_store(), *a)
    except Exception as exc:  # noqa: BLE001 - the agent should see why a build failed
        return f"ERROR building: {type(exc).__name__}: {exc}"


@tool
def build_fund_tables() -> str:
    """Build sections 2.1 (Fund Summary), 2.2 (LP Commitments), 4.1 (Balance Sheet) and
    4.2 (Statement of Operations) from the Delaware package, the Mauritius workpapers and
    the Fund Model. They are written straight into the report, every figure recorded and
    grounded, with prior-quarter comparatives. You do not type these numbers. You may add
    narrative to those sections afterwards with report_set_section (the tables stay)."""
    out = [_build(B.build_balance_sheet), _build(B.build_operations),
           _build(B.build_lp_commitments), _build(B.build_fund_summary)]
    return "\n".join(out)


@tool
def build_portfolio_tables() -> str:
    """Build sections 5.1 (sector table + chart), 5.2 (country table + chart), 5.3
    (composition chart), 5.4 (Portfolio Performance Summary) and 5.5 (Jobs & Impact),
    straight from the Fund Model and the Portfolio Metrics workbook. Every figure is
    recorded and grounded. Add narrative afterwards with report_set_section if wanted."""
    return "\n".join([_build(B.build_portfolio), _build(B.build_jobs)])


@tool
def ledger_search(text: str, limit: int = 25) -> str:
    """Find facts already in the ledger by part of their label, with value, unit and status.

    Use this BEFORE report_derive_fact so you use exact labels instead of guessing them, and
    to check whether a figure you want to quote is already grounded. Only 'extracted' and
    'derived' facts can be used in the report."""
    rows = get_store().find_facts(text, limit=min(int(limit), 60))
    if not rows:
        return f"no ledger facts match {text!r}"
    out = [f"{len(rows)} match(es) for {text!r} (label | value | unit | status):"]
    for r in rows:
        v = "null" if r["value"] is None else f"{r['value']:,.6g}"
        out.append(f"  {r['label']} | {v} | {r['unit']} | {r['status']}")
    return "\n".join(out)


@tool
def report_review_note(area: str, text: str, severity: str = "info") -> str:
    """Record something about the DATA for the human reviewer. It is NOT part of the report.

    Use this for every gap, conflict between sources, stale or carried-forward mark,
    unit problem, judgement call, or decision you need from the user. severity is
    'decision' (the user must choose), 'warning', or 'info'. The report itself must be a
    finished document with none of this in it: never write about missing data, sources,
    or how a figure was obtained inside a section, table or chart.
    """
    if severity not in {"decision", "warning", "info"}:
        return "ERROR: severity must be decision, warning or info."
    get_store().add_review_note(area.strip()[:60], text.strip(), severity)
    return f"review note recorded under '{area}' ({severity}); it will not appear in the report."


@tool
def report_set_cover(quarter: str = "", period: str = "", reporting_date: str = "",
                     jurisdiction: str = "Delaware / Mauritius", version: str = "v1.0") -> str:
    """Set the cover page fields. The title, fund name, logo, confidentiality line and
    design are fixed by the brand; only these values vary."""
    P = pr.current()
    quarter, period, reporting_date = quarter or P.label, period or P.range_label, reporting_date or P.end_label
    get_store().ensure_report("Chui Ventures Fund I", quarter)
    get_store().set_meta(quarter=quarter, period=period, reporting_date=reporting_date,
                         jurisdiction=jurisdiction, version=version)
    return f"cover set: {quarter}, {period}, reporting date {reporting_date}."


@tool
def report_remove_section(key: str) -> str:
    """Remove a section (and with it its tables and charts) from the report. Use it for any
    section the data does not let you write as finished content: omit it, and log why with
    report_review_note. A report shows only what it can stand behind."""
    st = get_store()
    st.delete_section(key)
    return f"section {key} removed from the report."


@tool
def inspect_pages(pages: str = "1", focus: str = "") -> str:
    """LOOK at rendered pages of the current PDF and report what is wrong visually.

    `pages` like "1", "1-3" or "2,5,9" (max 4 per call). `focus` says what to check, e.g.
    "is the logo placed correctly at top-left and fully visible", "does the table fit the
    page and are columns readable". Returns a plain critique from a vision model: layout,
    overlap, clipping, alignment, placement of the logo, empty pages, wrapped headers.
    Render first with report_render; then inspect; fix what you can; render again.
    """
    import base64
    import io

    import pypdfium2 as pdfium
    from langchain_core.messages import HumanMessage

    from ..render.report_writer import OUT_DIR
    from .llm import get_llm

    # The PDF of the latest render, exactly (not "the newest file in the folder"), read whole so a
    # concurrent re-render cannot change it under us.
    latest = (get_store().meta() or {}).get("last_render", {}).get("pdf")
    target = Path(latest) if latest and Path(latest).exists() else None
    if target is None:
        found = sorted(OUT_DIR.glob("*.pdf"), key=lambda f: f.stat().st_mtime, reverse=True)
        target = found[0] if found else None
    if target is None:
        return "ERROR: no rendered PDF yet; call report_render first."
    pdfs = [target]
    data = target.read_bytes()
    from ..render.pdfium_safe import PDFIUM_LOCK

    want: list[int] = []
    for part in pages.replace(" ", "").split(","):
        a, _, b = part.partition("-")
        want += list(range(int(a), int(b or a) + 1))
    images: list[str] = []
    with PDFIUM_LOCK:                       # PDFium crashes the process if two threads use it at once
        try:
            doc = pdfium.PdfDocument(data)
        except Exception as exc:  # noqa: BLE001
            return (f"ERROR: could not open {target.name} ({len(data)} bytes, starts {data[:8]!r}): {exc}. "
                    "Render again with report_render and retry once; if it still fails, record it with "
                    "report_review_note and carry on. This is not something to ask the user about.")
        try:
            total = len(doc)
            want = [p for p in want if 1 <= p <= total][:4]
            if not want:
                return f"ERROR: the PDF has {total} pages."
            for n in want:
                buf = io.BytesIO()
                doc[n - 1].render(scale=1.15).to_pil().convert("RGB").save(buf, "JPEG", quality=82)
                images.append("data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode())
        finally:
            doc.close()
    content = [{"type": "text", "text": (
        "You are the layout reviewer for a branded investor report (navy #00223B and tangerine #E25A00, "
        "Larken serif type). Look at each page image and report visual problems ONLY: text overlapping, "
        "clipped or overflowing content, tables running off the page or columns too narrow so words break "
        "badly, misaligned headers vs cells, the logo missing / cropped / badly placed / wrong size, large "
        "unexplained blank areas, headings orphaned at a page bottom, low-contrast text. "
        "Be specific about the page and the element. If a page looks right say 'Page N: fine'. "
        f"Pages shown, in order: {want}. " + (f"Focus especially on: {focus}" if focus else ""))}]
    for url in images:
        content.append({"type": "image_url", "image_url": {"url": url}})
    try:
        reply = get_llm().invoke([HumanMessage(content=content)])
    except Exception as exc:  # noqa: BLE001
        return f"ERROR: the vision review failed: {type(exc).__name__}: {exc}"
    get_store().mark("inspected_at")
    return f"[{pdfs[0].name}, {total} pages; reviewed {want}]\n" + str(reply.content)


@tool
def report_set_section(key: str, title: str, body_markdown: str, order: int = 0) -> str:
    """Create or replace a narrative section of the report.

    Sections are addressed by key (e.g. '1.1', '3.1') and can be rewritten
    independently -- this is how the report is mutated after it is first built.
    """
    get_store().ensure_report("Chui Ventures Fund I", pr.current().label)
    get_store().set_section(key, title, body_markdown, order)
    return f"section {key} saved ({len(body_markdown)} chars)"


@tool
def report_set_table(key: str, title: str, columns_json: str, rows_json: str,
                     section_key: str = "", options_json: str = "") -> str:
    """Create or replace a table in the report.

    `columns_json` is a JSON list of column headers; `rows_json` a JSON list of
    rows, each a list of cell values.
    """
    get_store().ensure_report("Chui Ventures Fund I", pr.current().label)
    try:
        cols = json.loads(columns_json)
        rows = json.loads(rows_json)
    except json.JSONDecodeError as exc:
        return f"ERROR: invalid JSON: {exc}"
    try:
        options = json.loads(options_json) if options_json else {}
    except json.JSONDecodeError as exc:
        return f"ERROR: options_json is not valid JSON: {exc}"
    get_store().set_table(key, title, cols, rows, section_key or None, options)
    return f"table {key} saved ({len(rows)} rows x {len(cols)} cols)"


@tool
def report_outline() -> str:
    """Show the report's current sections and tables."""
    secs = get_store().sections()
    tbls = get_store().tables()
    if not secs and not tbls:
        return "report is empty"
    out = ["SECTIONS:"]
    for s in secs:
        out.append(f"  {s['key']:6} {s['title'][:50]:52} {len(s['body'])} chars")
    out.append("TABLES:")
    for k, t in sorted(tbls.items()):
        out.append(f"  {k:6} {str(t['title'])[:50]:52} {len(t['rows'])} rows")
    out.append(f"facts in ledger: {get_store().fact_count()}")
    return "\n".join(out)


@tool
def report_render(file_stem: str = "Chui Ventures Fund I - Report") -> str:
    """Render the stored report to .docx and .pdf.

    Every number in narrative text is checked against the fact ledger first; the
    render is refused if an ungrounded figure is present.
    """
    from ..render.report_writer import render_report

    try:
        docx_path, pdf_path, report = render_report(get_store(), file_stem)
    except Exception as exc:  # noqa: BLE001
        return f"RENDER BLOCKED: {exc}"
    get_store().set_meta(last_render={"docx": str(docx_path), "pdf": str(pdf_path),
                                      "notes": report["review_notes"], "headings": report["heading_pages"]})
    get_store().mark("rendered_at")
    return (f"rendered {report['sections']} sections, {report['tables']} tables\n"
            f"docx: {docx_path}\npdf:  {pdf_path}")


# --------------------------------------------------------------------------


def _resolve_typed(name: str, suffixes: tuple[str, ...]) -> Path:
    """Prefer files of the right kind; if the name only matches another kind, return that so
    the tool can say 'this is a PDF, use read_pdf' rather than 'not found'."""
    try:
        return _resolve(name, suffixes)
    except FileNotFoundError:
        return _resolve(name)


def _resolve(name: str, suffixes: tuple[str, ...] | None = None) -> Path:
    """Resolve a partial file name against the source tree.

    A tool that only reads one kind of file passes `suffixes`, so "Uncover" finds the
    drawdown PDF for read_pdf and the valuation workbook for the Excel tools."""
    name_l = name.casefold()
    exact = config.SOURCE_ROOT / name
    if exact.exists() and (not suffixes or exact.suffix.lower() in suffixes):
        return exact
    matches = [p for p in config.SOURCE_ROOT.rglob("*")
               if p.is_file() and name_l in p.name.casefold() and not p.name.startswith("~$")
               and "chui-reporter" not in p.parts and ".venv" not in p.parts
               and (not suffixes or p.suffix.lower() in suffixes)]
    if not matches:
        kind = f" ({'/'.join(suffixes)})" if suffixes else ""
        raise FileNotFoundError(f"no source file matching {name!r}{kind}")
    return matches[0]


def _fmt(v) -> str:
    if v is None:
        return "null"
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)


# --------------------------------------------------------------------------
# Asking the user
# --------------------------------------------------------------------------


def _budget_left(call_id: str) -> bool:
    if call_id in _asked:                 # the same call replayed after the user answered
        return True
    if len(_asked) >= QUESTION_BUDGET:
        return False
    _asked.add(call_id)
    return True


@tool
def ask_user(question: str, why_it_matters: str, options: list[str] | None = None,
             tool_call_id: Annotated[str, InjectedToolCallId] = "") -> str:
    """Ask the user a question and wait for the answer. The run pauses until they reply.

    Use this ONLY when a person could answer and the sources cannot: a judgement the data
    does not settle (which of two conflicting figures is authoritative), or a fact that exists
    only in someone's head (what the GP decided about a governance matter). Never ask for
    something you can read from the files, never ask what you could settle by stating an
    assumption in a review note, and never ask more than once about the same thing.

    question: one plain sentence, written to the person reading it, no jargon.
    why_it_matters: one sentence on what in the report depends on the answer.
    options: up to three short, mutually exclusive answers the user can pick; omit when the
    answer is free text.

    Not for problems with your own tools: a failed render, an unreadable file, a tool error. Those are
    yours to retry, work around, or record with report_review_note."""
    if len(question.strip()) > 220 or len(why_it_matters.strip()) > 320:
        return ("ERROR: too long. The question is shown as a headline: ask it in one short sentence "
                "(under 200 characters) and give the context in why_it_matters (under 300 characters).")
    if not _budget_left(tool_call_id):
        return ("You have used your questions for this run. Decide yourself, state the assumption "
                "with report_review_note, and carry on.")
    answer = interrupt({"kind": "info", "prompt": question.strip(), "why": why_it_matters.strip(),
                        "options": [str(o)[:140] for o in (options or [])][:3]})
    return f"The user answered: {answer}"


@tool
def request_sources(slots: list[str], why_it_matters: str,
                    tool_call_id: Annotated[str, InjectedToolCallId] = "") -> str:
    """Tell the user which source documents are missing and wait until they add them.

    `slots` are ids from list_sources' kinds of source: delaware, lp_workpaper, fund_model,
    portfolio_metrics, valuations, prior_report, pipeline, macro, gp_statement. The run pauses;
    it resumes by itself when files covering the slots are added to the folder, or when the user
    chooses to continue without them. Ask once per gap, only for a source that would change what
    the report says, and carry on without it if the user declines."""
    from ..workspace.slots import BY_ID, slot_files

    unknown = [x for x in slots if x not in BY_ID]
    if unknown:
        return f"unknown source ids {unknown}; use ids from: {', '.join(BY_ID)}"
    paths = [p.relative_to(config.SOURCE_ROOT).as_posix() for p in config.SOURCE_ROOT.rglob("*")
             if p.is_file()]
    missing = [x for x in slots if len(slot_files(BY_ID[x], paths)) < BY_ID[x].minimum]
    if not missing:
        return "Those sources are already available; read them."
    if not _budget_left(tool_call_id):
        return ("You have used your questions for this run. Carry on without those sources, leave the "
                "sections that depend on them out, and record the gap with report_review_note.")
    answer = interrupt({"kind": "sources", "slots": missing, "why": why_it_matters.strip(),
                        "prompt": "Add " + ", ".join(BY_ID[x].label for x in missing) + " to your folder."})
    return f"{answer}"



ALL_TOOLS = [
    list_sources,
    ask_user,
    request_sources,
    read_text,
    look_at_image,
    read_pdf,
    excel_sheets,
    excel_find_value,
    excel_dump_region,
    portfolio_valuations,
    fund_capital_position,
    financial_statements,
    prior_report_table,
    report_save_facts,
    report_derive_fact,
    build_fund_tables,
    build_portfolio_tables,
    ledger_search,
    report_review_note,
    report_set_cover,
    report_remove_section,
    inspect_pages,
    report_set_section,
    report_set_table,
    report_outline,
    report_render,
]
