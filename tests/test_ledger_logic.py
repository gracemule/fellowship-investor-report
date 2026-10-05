"""How a number becomes allowed in the report, tested on a small synthetic workbook.

The headline test is `test_a_fabricated_source_cannot_ground_a_number`: in the first design the
agent wrote the ledger itself, so an invented figure saved with an invented source passed the gate.
(The same safeguards are exercised against the real documents in tests/private, which is not published.)
"""

from __future__ import annotations

import json

import pytest
from openpyxl import Workbook as XlWorkbook

from chui_reporter import config
from chui_reporter.agent import tools as T
from chui_reporter.agent.ledger import DerivationError, evaluate
from chui_reporter.agent.store import Fact
from chui_reporter.render.gate import check_grounded


@pytest.fixture()
def tools_store(store, monkeypatch, tmp_path):
    """Tools pointed at an isolated schema and a synthetic source folder."""
    monkeypatch.setattr(T, "_STORE", store)
    wb = XlWorkbook()
    ws = wb.active
    ws.title = "Balances"
    ws["A5"], ws["C5"] = "Contributed capital", 4_218_640.55
    ws["A6"], ws["C6"] = "Notes", "Reviewed by the administrator"
    (tmp_path / "Fund Financials").mkdir()
    wb.save(tmp_path / "Fund Financials" / "synthetic-workpaper.xlsx")
    old = config.SOURCE_ROOT
    config.set_root(tmp_path)
    yield store
    config.set_root(old)


# -- derivation is arithmetic, nothing else ----------------------------------


def test_evaluate_does_arithmetic_over_named_inputs():
    assert evaluate("a - b", {"a": 614_545.27, "b": 940_000.0}) == pytest.approx(-325_454.73)
    assert evaluate("(a + b) / 2", {"a": 10, "b": 20}) == 15
    assert evaluate("-a", {"a": 3}) == -3


@pytest.mark.parametrize("expr", [
    "__import__('os').system('x')", "open('f')", "a.real", "a ** 2", "a if b else a",
    "[a]", "'text'", "lambda: 1",
])
def test_evaluate_refuses_anything_but_arithmetic(expr):
    with pytest.raises(DerivationError):
        evaluate(expr, {"a": 1, "b": 2})


def test_evaluate_refuses_unknown_names_and_zero_division():
    with pytest.raises(DerivationError, match="unknown input"):
        evaluate("a + c", {"a": 1})
    with pytest.raises(DerivationError, match="zero"):
        evaluate("a / b", {"a": 1, "b": 0})


# -- the ledger cannot be self-certified ---------------------------------------


def _save(facts):
    return T.report_save_facts.invoke({"facts_json": json.dumps(facts)})


def test_a_correctly_cited_claim_licenses_the_figure(tools_store):
    _save([{"label": "contributed", "value": 4_218_640.55, "unit": "USD",
            "source_file": "synthetic-workpaper.xlsx", "source_sheet": "Balances", "source_cell": "C5"}])
    assert check_grounded("Contributed capital was $4,218,640.55.", tools_store.grounded_values()) == []


def test_a_fabricated_source_cannot_ground_a_number(tools_store):
    """The attack: save an invented figure citing a real file and an invented cell."""
    out = _save([{"label": "made up", "value": 9_800_000, "unit": "USD",
                  "source_file": "synthetic-workpaper.xlsx", "source_sheet": "Balances", "source_cell": "C99"}])
    assert "0 verified" in out and "CLAIMED" in out
    assert 9_800_000.0 not in tools_store.grounded_values()
    assert check_grounded("Fair value was $9.8M.", tools_store.grounded_values()) == ["$9.8M"]


def test_a_real_cell_with_the_wrong_value_is_rejected(tools_store):
    out = _save([{"label": "wrong", "value": 4_000_000, "unit": "USD",
                  "source_file": "synthetic-workpaper.xlsx", "source_sheet": "Balances", "source_cell": "C5"}])
    assert "CLAIMED" in out and 4_000_000.0 not in tools_store.grounded_values()


def test_a_claimed_fact_cannot_be_a_derivation_input(tools_store):
    tools_store.add_facts([Fact("unverified", 100.0, status="claimed"),
                           Fact("good", 40.0, source_file="a.xlsx", status="extracted")])
    out = T.report_derive_fact.invoke({"label": "d", "expression": "a - b",
                                       "inputs_json": json.dumps({"a": "unverified", "b": "good"})})
    assert out.startswith("ERROR") and "not grounded" in out


def test_derived_figures_license_prose_but_are_computed_not_typed(tools_store):
    tools_store.add_facts([Fact("a", 614_545.45, source_file="x.xlsx", status="extracted"),
                           Fact("b", 940_000.0, source_file="y.xlsx", status="extracted")])
    out = T.report_derive_fact.invoke({"label": "gap", "expression": "a - b",
                                       "inputs_json": json.dumps({"a": "a", "b": "b"})})
    assert "derived and recorded" in out
    assert check_grounded("a gap of $325,454.55", tools_store.grounded_values()) == []


def test_a_label_that_swallowed_the_call_arguments_is_rejected(tools_store):
    """A malformed call once produced a ledger label containing a chunk of JSON."""
    bad = 'Uncalled (derived)","expression":"a - b","inputs_json":"{}"'
    out = T.report_derive_fact.invoke({"label": bad, "expression": "a - b",
                                       "inputs_json": json.dumps({"a": "x", "b": "y"})})
    assert out.startswith("ERROR: label") and tools_store.fact_count() == 0
