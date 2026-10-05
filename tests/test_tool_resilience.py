"""Regressions from the first live run, where a tool crash killed the whole agent.

The agent called an Excel tool on a PDF; openpyxl raised; nothing caught it; the
run died with two tool calls still pending. A tool failure must be information
the agent can act on, never the end of the run.
"""

from __future__ import annotations

import pytest
from langchain_core.messages import AIMessage
from langchain_core.tools import tool
from langgraph.prebuilt import ToolNode

from chui_reporter.agent import tools as T


def test_a_raising_tool_becomes_a_tool_message_not_a_crash():
    @tool
    def explode(x: str) -> str:
        """Always fails."""
        raise RuntimeError("boom")

    from langgraph.graph import END, START, MessagesState, StateGraph

    g = StateGraph(MessagesState)
    g.add_node("tools", ToolNode([explode], handle_tool_errors=True))
    g.add_edge(START, "tools")
    g.add_edge("tools", END)
    msg = AIMessage(content="", tool_calls=[{"name": "explode", "args": {"x": "1"},
                                              "id": "c1", "type": "tool_call"}])
    out = g.compile().invoke({"messages": [msg]})["messages"][-1]
    assert out.status == "error" and "boom" in out.content


def test_the_agent_graph_tolerates_tool_errors(monkeypatch):
    """Pin the actual wiring: the graph's tool node must swallow tool exceptions."""
    from chui_reporter.agent.graph import build_agent

    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-offline")
    node = build_agent(None, "deepseek").nodes["tools"].bound
    assert node._handle_tool_errors is True
    assert "read_pdf" in node.tools_by_name


# -- fact validation: found when Postgres rejected a date in a numeric column --


def test_as_number_accepts_figures_and_rejects_dates_and_text():
    assert T._as_number(5) == 5.0
    assert T._as_number("1,234.50") == 1234.5
    assert T._as_number(None) is None and T._as_number("") is None
    for bad in ("2026-10-01", "n/a", "$9.7M", True, float("nan"), float("inf")):
        with pytest.raises(ValueError):
            T._as_number(bad)


def test_save_facts_rejects_a_date_as_value_without_saving_anything(monkeypatch):
    saved = []

    class Spy:
        def add_facts(self, f):
            saved.extend(f)
            return len(f)

        def fact_count(self):
            return len(saved)

    monkeypatch.setattr(T, "get_store", lambda: Spy())
    import json

    out = T.report_save_facts.invoke({"facts_json": json.dumps([
        {"label": "good", "value": 100, "source_file": "a.xlsx"},
        {"label": "drawdown_date", "value": "2026-10-01", "source_file": "d.pdf"},
    ])})
    assert out.startswith("ERROR: nothing saved") and "drawdown_date" in out
    assert "text_value" in out, "the message must tell the agent how to fix it"
    assert saved == [], "one bad fact must not save the good ones half-way"




# -- read_text: the formats that arrive without a dedicated reader ----------------------------


def test_read_text_reads_text_csv_and_word_and_points_other_kinds_elsewhere(tmp_path):
    from docx import Document

    from chui_reporter import config

    (tmp_path / "Macro and Context").mkdir()
    (tmp_path / "Macro and Context" / "notes.txt").write_text("Policy rate held.\n")
    (tmp_path / "Macro and Context" / "gdp.csv").write_text("country,gdp\nKenya,5.1\n")
    d = Document()
    d.add_paragraph("GP statement")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Matter", "Resolved"
    d.save(tmp_path / "Macro and Context" / "gp.docx")
    old = config.SOURCE_ROOT
    config.set_root(tmp_path)
    try:
        assert "Policy rate held." in T.read_text.invoke({"file_name": "notes"})
        assert "Kenya,5.1" in T.read_text.invoke({"file_name": "gdp.csv"})
        out = T.read_text.invoke({"file_name": "gp.docx"})
        assert "GP statement" in out and "Matter | Resolved" in out
        assert T.read_text.invoke({"file_name": "nothing-like-this"}).startswith("ERROR")
        cut = T.read_text.invoke({"file_name": "notes", "max_chars": 500})
        assert "more characters" not in cut
    finally:
        config.set_root(old)
