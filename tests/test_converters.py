"""The PDF converter layer: a hosted converter is exercised against a stand-in for the service (no
network, no credit), and the Word file's font embedding is checked for validity."""

from __future__ import annotations

import io
import uuid
import zipfile
from pathlib import Path

import httpx
import pytest
from docx import Document

from chui_reporter.render import converters as C
from chui_reporter.render.convert import ConversionError

PDF = b"%PDF-1.7\n%stand-in\n"


def _docx(tmp_path: Path, text="Hello") -> Path:
    d = Document()
    d.add_paragraph(text).runs[0].font.name = "Larken-Regular"
    p = tmp_path / "report.docx"
    d.save(str(p))
    return p


class FakeILove:
    """Just enough of the service: auth, start, upload, process, download, delete."""

    def __init__(self, *, quota=False, zipped=False, bad_token_first=False, start_method="GET"):
        self.calls, self.quota, self.zipped = [], quota, zipped
        self.bad_token_first, self.start_method = bad_token_first, start_method
        self.uploaded = b""

    def __call__(self, req: httpx.Request) -> httpx.Response:
        path = req.url.path
        self.calls.append(f"{req.method} {path}")
        if path == "/v1/auth":
            return httpx.Response(200, json={"token": "tok-1"})
        if self.bad_token_first and req.headers.get("authorization") == "Bearer tok-1" and not any("retried" in c for c in self.calls):
            self.calls.append("retried")
            return httpx.Response(401)
        if path.startswith("/v1/start/officepdf"):
            if self.quota:
                return httpx.Response(403, json={"error": {"message": "Your monthly credit limit has been reached"}})
            if req.method != self.start_method:
                return httpx.Response(405)
            return httpx.Response(200, json={"server": "api9.ilovepdf.com", "task": "T1", "remaining_credits": 233})
        if path == "/v1/upload":
            self.uploaded = req.content
            return httpx.Response(200, json={"server_filename": "srv.docx"})
        if path == "/v1/process":
            return httpx.Response(200, json={"status": "TaskSuccess", "download_filename": "report.pdf"})
        if path.startswith("/v1/download/"):
            if self.zipped:
                buf = io.BytesIO()
                with zipfile.ZipFile(buf, "w") as z:
                    z.writestr("report.pdf", PDF)
                return httpx.Response(200, content=buf.getvalue())
            return httpx.Response(200, content=PDF)
        if path.startswith("/v1/task/"):
            return httpx.Response(200, json={})
        return httpx.Response(404)


def _conv(fake, tmp_path, **kw):
    c = C.ILoveApiConverter("pub-key", client=httpx.Client(transport=httpx.MockTransport(fake)), embed=False,
                            cache_dir=tmp_path / "cache", sleep=lambda s: None, **kw)
    return c


def test_the_hosted_flow_runs_in_order_and_cleans_up(tmp_path):
    fake = FakeILove()
    out = _conv(fake, tmp_path).convert(_docx(tmp_path), tmp_path / "out")
    assert out.read_bytes() == PDF and out.name == "report.pdf"
    order = [c for c in fake.calls]
    assert order[:5] == ["POST /v1/auth", "GET /v1/start/officepdf", "POST /v1/upload", "POST /v1/process", "GET /v1/download/T1"]
    assert order[-1] == "DELETE /v1/task/T1", "the unfinished report is not left on their server"


def test_an_identical_document_is_never_converted_or_paid_for_twice(tmp_path):
    fake, used = FakeILove(), []
    c = _conv(fake, tmp_path, on_use=lambda left: used.append(left))
    doc = _docx(tmp_path)
    c.convert(doc, tmp_path / "out"); c.convert(doc, tmp_path / "out")
    assert sum(1 for x in fake.calls if x.startswith("POST /v1/upload")) == 1
    assert used.count(None) == 1 and 233 in used, "one conversion counted; the service's own credit count is passed on"
    c.convert(_docx(tmp_path, "different words"), tmp_path / "out")
    assert sum(1 for x in fake.calls if x.startswith("POST /v1/upload")) == 2


def test_a_used_up_allowance_is_reported_as_such(tmp_path):
    with pytest.raises(C.ConverterQuota, match="allowance is used up"):
        _conv(FakeILove(quota=True), tmp_path).convert(_docx(tmp_path), tmp_path / "out")


def test_an_expired_token_is_replaced_and_the_call_repeated(tmp_path):
    out = _conv(FakeILove(bad_token_first=True), tmp_path).convert(_docx(tmp_path), tmp_path / "out")
    assert out.read_bytes() == PDF


def test_zipped_output_and_a_post_only_start_are_both_handled(tmp_path):
    out = _conv(FakeILove(zipped=True, start_method="POST"), tmp_path).convert(_docx(tmp_path), tmp_path / "out")
    assert out.read_bytes() == PDF


def test_a_response_that_is_not_a_pdf_is_an_error_not_a_report(tmp_path):
    class Junk(FakeILove):
        def __call__(self, req):
            if req.url.path.startswith("/v1/download/"):
                return httpx.Response(200, content=b"<html>error</html>")
            return super().__call__(req)

    with pytest.raises(ConversionError, match="did not return a PDF"):
        _conv(Junk(), tmp_path).convert(_docx(tmp_path), tmp_path / "out")


# -- font embedding ------------------------------------------------------------------------------


def test_obfuscation_is_its_own_inverse_and_changes_only_the_first_32_bytes():
    font, guid = bytes(range(256)) * 4, uuid.uuid4()
    once = C.obfuscate(font, guid)
    assert once[32:] == font[32:] and once[:32] != font[:32]
    assert C.obfuscate(once, guid) == font


def test_a_font_is_embedded_into_a_valid_word_file_and_the_original_is_untouched(tmp_path):
    src = _docx(tmp_path)
    before = src.read_bytes()
    dst = C.embed_fonts(src, tmp_path / "sent.docx", {"Larken-Regular": b"\x00\x01\x00\x00" + b"x" * 100,
                                                     "Larken-Bold": b"\x00\x01\x00\x00" + b"y" * 100})
    assert src.read_bytes() == before, "the Word file the user downloads is never embedded"
    with zipfile.ZipFile(dst) as z:
        names = z.namelist()
        assert names[0] == "[Content_Types].xml"
        assert {"word/fonts/font1.odttf", "word/fonts/font2.odttf", "word/_rels/fontTable.xml.rels"} <= set(names)
        table = z.read("word/fontTable.xml").decode()
        assert table.count("<w:embedRegular") == 2 and 'w:name="Larken-Bold"' in table
        assert "odttf" in z.read("[Content_Types].xml").decode()
    assert Document(str(dst)).paragraphs[0].text == "Hello", "still opens as a Word document"


def test_the_converter_choice_follows_the_environment(monkeypatch):
    monkeypatch.setenv("CHUI_PDF_CONVERTER", "iloveapi")
    monkeypatch.setenv("ILOVEAPI_PUBLIC_KEY", "k")
    assert isinstance(C.get_converter(), C.ILoveApiConverter)
    monkeypatch.setenv("CHUI_PDF_CONVERTER", "libreoffice")
    assert isinstance(C.get_converter(), C.LibreOfficeConverter)
    monkeypatch.setenv("CHUI_PDF_CONVERTER", "auto")
    monkeypatch.setattr(C, "find_soffice", lambda: (_ for _ in ()).throw(ConversionError("none")))
    assert isinstance(C.get_converter(), C.ILoveApiConverter), "no LibreOffice, but a key: use the hosted one"
    monkeypatch.delenv("ILOVEAPI_PUBLIC_KEY")
    with pytest.raises(ConversionError, match="no PDF converter"):
        C.get_converter()


def test_a_hosted_converter_without_the_font_files_says_what_to_add(tmp_path, monkeypatch):
    monkeypatch.setattr(C, "font_files", lambda root=None: {})
    c = C.ILoveApiConverter("k", embed=True)
    with pytest.raises(ConversionError, match="Larken font files are missing"):
        c.prepare(tmp_path)


def test_font_names_are_read_out_of_a_pdf():
    from pypdf import PdfWriter

    from chui_reporter.render.check_converter import pdf_fonts

    w = PdfWriter()
    w.add_blank_page(100, 100)
    buf = io.BytesIO()
    w.write(buf)
    assert pdf_fonts(buf.getvalue()) == []


@pytest.mark.slow
def test_libreoffice_accepts_a_word_file_with_embedded_fonts(tmp_path):
    from chui_reporter.render.convert import docx_to_pdf

    sent = C.embed_fonts(_docx(tmp_path), tmp_path / "e.docx", {"Larken-Regular": (C.font_files() or {"Larken-Regular": b"x"})["Larken-Regular"]})
    assert docx_to_pdf(sent, tmp_path).exists()
