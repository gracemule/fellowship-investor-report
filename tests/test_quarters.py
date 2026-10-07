"""Moving through the quarters: each one is a clean slate with files, runs, notes and conversation of its own, and only the
brand kit (and the previous quarter's report, on request) carries over."""

from __future__ import annotations

import hashlib
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from langgraph.checkpoint.memory import InMemorySaver

from chui_reporter.agent.store import Fact, Store
from chui_reporter.app.auth import Auth
from chui_reporter.app.main import create_app
from chui_reporter.runtime import state
from chui_reporter.runtime.runner import Runtime
from chui_reporter.workspace import sync
from chui_reporter.workspace.slots import is_durable

Q2, Q3 = "2026Q2", "2026Q3"


def _h(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def _put(store, ws, path, data=b"x"):
    sync.put_file(store, path, data, _h(data), workspace=ws)
    return {"path": path, "sha256": _h(data)}


def _paths(store, ws):
    return [r["path"] for r in sync._present(store, ws)]


FONTS = [f"Branding/Fonts/Larken/Larken {w}.ttf" for w in ("Regular", "Bold", "Italic")]
LOGO = "Branding/cv_logo_white.png"


def _brand(store, ws):
    return [_put(store, ws, p, b"font:" + p.encode()) for p in FONTS] + [_put(store, ws, LOGO, b"png")]


@pytest.fixture()
def rt(store):
    r = Runtime(store, saver=InMemorySaver(), agent_factory=lambda p: None, prepare=False, snapshot=False, sleep=lambda s: None)
    state.set_session_provider(lambda: r.session()["id"])
    yield r
    state.set_session_provider(None)


# ---- files --------------------------------------------------------------------------------------------------------


def test_a_blank_folder_for_one_quarter_never_touches_another_quarters_files(store):
    sync.commit(store, [_put(store, Q2, "Portfolio Company Data/Fund Model.xlsx"), _put(store, Q2, "Fund Financials/LP.xlsx")], Q2)
    assert len(_paths(store, Q2)) == 2
    sync.commit(store, [], Q3, folder="q3 data")                       # the user attached an empty folder for Q3
    assert len(_paths(store, Q2)) == 2, "Q2 is exactly as it was"
    assert _paths(store, Q3) == []


def test_the_brand_kit_is_stored_once_and_serves_every_quarter(store):
    sync.commit(store, _brand(store, Q2), Q2)
    assert _paths(store, Q2) == [] and len(_paths(store, sync.SHARED)) == 4, "the brand kit is not filed under a quarter"
    assert all(is_durable(p) for p in FONTS + [LOGO])
    cov = {c.slot.id: c for c in sync.coverage(store, workspace=Q3)}
    assert cov["brand_fonts"].state == "ready" and cov["brand_logos"].state == "ready", "Q3 never asks for it again"
    sync.commit(store, [], Q3)                                          # a Q3 folder without a Branding folder
    assert len(_paths(store, sync.SHARED)) == 4, "a folder that lacks the brand kit does not remove it"
    assert {c.slot.id for c in sync.required_missing(sync.coverage(store, workspace=Q3))} >= {"delaware", "fund_model"}


def test_a_changed_brand_file_is_a_change_and_an_unchanged_one_is_not(store):
    m = _brand(store, Q2)
    assert sorted(sync.commit(store, m, Q2).added) == sorted(FONTS + [LOGO])
    assert not sync.commit(store, m, Q3).any, "the same brand kit seen again under another quarter changes nothing"
    new = [_put(store, Q3, LOGO, b"a new logo")]
    assert sync.commit(store, new, Q3).modified == [LOGO]


def test_clearing_a_quarter_removes_its_synced_files_but_not_the_brand_kit_or_attachments(store):
    sync.commit(store, _brand(store, Q3) + [_put(store, Q3, "Fund Financials/LP.xlsx")], Q3, folder="wrong folder")
    sync.put_attachment(store, "note.csv", b"a,b\n", Q3)
    assert sync.clear_quarter(store, Q3) == 1
    assert _paths(store, Q3) == ["Uploads/note.csv"] and len(_paths(store, sync.SHARED)) == 4
    ws = state.workspace(store, Q3)
    assert "folder" not in (ws["settings"] or {}) and ws["last_sync_at"] is None


def test_moving_a_quarters_files_replaces_what_the_target_already_has(store):
    sync.commit(store, [_put(store, Q3, "Fund Financials/LP.xlsx", b"q2 numbers"), _put(store, Q3, "Valuation Reports/A.xlsx", b"a")],
                Q3, folder="my folder")
    _put(store, Q2, "Fund Financials/LP.xlsx", b"old")
    assert sync.move_quarter(store, Q3, Q2) == 2
    assert _paths(store, Q3) == [] and sorted(_paths(store, Q2)) == ["Fund Financials/LP.xlsx", "Valuation Reports/A.xlsx"]
    assert state.workspace(store, Q2)["settings"]["folder"] == "my folder"
    assert sync.move_quarter(store, Q3, Q2) == 0 and sync.move_quarter(store, Q2, Q2) == 0


def test_materialize_writes_the_quarter_the_brand_kit_and_the_baseline_and_nothing_else(store, tmp_path):
    sync.commit(store, _brand(store, Q3) + [_put(store, Q3, "Fund Financials/LP.xlsx", b"lp")], Q3)
    _put(store, Q2, "Fund Financials/Other quarter.xlsx", b"must not appear")
    dest = tmp_path / "src"
    (dest).mkdir()
    (dest / "stale.txt").write_text("left over")
    n = sync.materialize(store, dest, Q3, extra={"Prior Period Baseline/Fund - Q2 2026 Investor Report.pdf": b"%PDF-prior"})
    assert n == 6
    got = sorted(p.relative_to(dest).as_posix() for p in dest.rglob("*") if p.is_file())
    assert got == sorted(FONTS + [LOGO, "Fund Financials/LP.xlsx", "Prior Period Baseline/Fund - Q2 2026 Investor Report.pdf"])
    assert sync.materialize(store, dest, Q3, extra={"Prior Period Baseline/Fund - Q2 2026 Investor Report.pdf": b"%PDF-prior"}) == 0


def test_the_previous_quarters_report_stands_in_for_the_baseline_until_the_user_supplies_one():
    paths = ["Fund Financials/LP.xlsx"]
    prior = {"label": "Q2 2026", "version": 8}
    c = {x.slot.id: x for x in sync.coverage_of(paths, prior=prior)}["prior_report"]
    assert c.state == "ready" and "Q2 2026 report built here (version 8)" in c.note and c.files == []
    assert {x.slot.id: x for x in sync.coverage_of(paths)}["prior_report"].state == "missing"
    own = {x.slot.id: x for x in sync.coverage_of(["Prior Period Baseline/Q2 report.pdf"], prior=prior)}["prior_report"]
    assert own.state == "ready" and own.note == "" and own.files, "a file the user supplied is what is shown"


# ---- the quarter as the page sees it -----------------------------------------------------------------------------------


def _build_q2(rt, store):
    q2 = Store(url=store.url, schema=store.schema, report_id=f"fund-i-{Q2}")
    q2.ensure_report("Chui Ventures Fund I", "Q2 2026")
    state.save_version(q2, 14, {"sections": []}, b"%PDF-q2", b"docx", None)
    return q2


def test_switching_quarter_is_a_clean_slate_and_switching_back_restores_the_old_one(store, rt):
    rt.set_period(Q2)
    sync.commit(store, _brand(store, Q2) + [_put(store, Q2, "Fund Financials/LP.xlsx")], Q2, folder="q2 data")
    q2 = _build_q2(rt, store)
    rid = rt._create("build", "Build it")
    state.update_run(store, rid, status="failed", error="boom")
    S2 = rt.snapshot_state()
    assert S2["workspace"]["files"] == 1 and S2["workspace"]["folder"] == "q2 data" and S2["version"]["version"] == 1
    assert S2["run"]["id"] == rid

    S3 = rt.set_period(Q3)
    assert S3["period"]["label"] == "Q3 2026" and S3["period"]["prev"] == "Q2 2026"
    assert S3["version"] is None and S3["versions"] == [] and S3["run"] is None, "no report, no old run"
    assert S3["workspace"]["files"] == 0 and S3["workspace"]["folder"] is None
    assert S3["status"]["phase"] == "empty" and "Q3 2026" in S3["status"]["headline"]
    cov = {c["id"]: c for c in S3["coverage"]}
    assert cov["brand_fonts"]["state"] == "ready" and cov["brand_fonts"]["durable"] is True, "the brand kit stands"
    assert cov["prior_report"]["state"] == "ready" and "Q2 2026" in cov["prior_report"]["note"], "the Q2 report is the baseline"
    assert cov["delaware"]["state"] == "missing"
    assert rt.report_store().report_id == f"fund-i-{Q3}" and state.latest_version(rt.report_store()) is None
    assert rt.session()["period"] == Q3 and rt.thread_id() != q2.report_id

    back = rt.set_period(Q2)
    assert back["workspace"]["files"] == 1 and back["version"]["version"] == 1 and back["run"]["id"] == rid


def test_a_synced_but_empty_folder_is_reported_as_connected_and_empty_not_as_nothing_chosen(store, rt):
    rt.set_period(Q3)
    assert rt.snapshot_state()["status"]["phase"] == "empty"
    sync.commit(store, [], Q3, folder="q3 data")
    s = rt.snapshot_state()
    assert s["status"]["phase"] == "empty_folder" and "q3 data" in s["status"]["headline"] and s["workspace"]["folder"] == "q3 data"


def test_the_quarter_cannot_change_while_the_agent_is_working(store, rt):
    rt.set_period(Q2)
    rid = rt._create("build", "Build it")
    state.update_run(store, rid, status="running")
    with pytest.raises(RuntimeError, match="busy"):
        rt.set_period(Q3)
    assert rt.period().code == Q2


def test_the_previous_report_becomes_the_baseline_file_unless_the_user_has_supplied_one(store, rt):
    rt.set_period(Q2)
    _build_q2(rt, store)
    rt.set_period(Q3)
    info = rt.prior_info()
    assert info["label"] == "Q2 2026" and info["version"] == 1
    assert list(rt.prior_baseline()) == ["Prior Period Baseline/Chui Ventures Fund I - Q2 2026 Investor Report.pdf"]
    assert list(rt.prior_baseline().values()) == [b"%PDF-q2"]
    _put(store, Q3, "Prior Period Baseline/Q2 2026 Report.pdf", b"theirs")
    assert rt.prior_baseline() == {}, "their own file wins"


def test_a_message_is_welcome_before_any_report_exists(store, rt):
    rt.set_period(Q3)
    r = rt.steer("hey")
    assert r["ok"] is True and r["applied"] == "new_run"
    run = state.get_run(store, r["run"])
    assert run["kind"] == "steer" and run["period"] == Q3


# ---- events keep the whole of what was said ----------------------------------------------------------------------------


def test_a_long_message_is_stored_whole_and_only_its_summary_is_short(store):
    long = ("Kenya inflation 6.41%, policy 8.75%. " * 60).strip()
    state.emit(store, "note", long)
    ev = state.recent_events(store, 5)[-1]
    assert len(ev["label"]) == state.LABEL_MAX and ev["detail"]["text"] == long
    state.emit(store, "note", "short")
    assert "text" not in state.recent_events(store, 5)[-1]["detail"]


# ---- review notes --------------------------------------------------------------------------------------------------------


def test_a_review_note_can_be_answered_resolved_and_reopened(store):
    rs = Store(url=store.url, schema=store.schema, report_id=f"fund-i-{Q3}")
    rs.ensure_report("Chui Ventures Fund I", "Q3 2026")
    rs.add_review_note("Macro snapshot", "Kenya: gdp growth is a dash in the table.", "warning")
    rs.add_review_note("Valuation", "Two marks are stale.", "info")
    first = rs.review_notes()[0]
    assert first["status"] == "open" and first["id"]
    assert rs.answer_review_note(first["id"], "KNBS says 5.3%")["status"] == "answered"
    assert rs.get_review_note(first["id"])["reply"] == "KNBS says 5.3%"
    res = rs.resolve_review_note(first["id"], "Used the figure you gave.")
    assert res["status"] == "resolved" and res["resolution"] == "Used the figure you gave." and res["resolved_at"]
    assert rs.answer_review_note(first["id"], "again") is None, "a resolved note is not answered again"
    assert [n["status"] for n in rs.review_notes()] == ["open", "resolved"], "open first, resolved last"
    assert rs.reopen_review_note(first["id"])["status"] == "open"
    assert rs.resolve_review_note(99999) is None


def test_notes_the_code_derives_do_not_bring_back_what_the_user_resolved(store):
    rs = Store(url=store.url, schema=store.schema, report_id=f"fund-i-{Q3}")
    rs.ensure_report("Chui Ventures Fund I", "Q3 2026")
    text = "Kenya: gdp growth is a dash in the table. KNBS unreachable."
    rs.replace_review_notes("Macro snapshot", [(text, "info")], area_like="macro")
    nid = rs.review_notes()[0]["id"]
    rs.resolve_review_note(nid, "Not needed this quarter")
    rs.replace_review_notes("Macro snapshot", [(text, "info"), ("Senegal: inflation is a dash.", "info")], area_like="macro")
    notes = {n["text"]: n["status"] for n in rs.review_notes()}
    assert notes[text] == "resolved" and notes["Senegal: inflation is a dash."] == "open" and len(notes) == 2


def test_replying_to_a_note_marks_it_answered_and_tells_the_agent_which_note(store, rt):
    rt.set_period(Q3)
    rs = rt.report_store()
    rs.add_review_note("Macro snapshot", "Kenya GDP is missing.", "warning")
    nid = rs.review_notes()[0]["id"]
    r = rt.reply_to_note(nid, "It was 5.3% in Q1 2026.")
    assert r["ok"] and rs.get_review_note(nid)["status"] == "answered"
    run = state.get_run(store, r["run"])
    assert f"review note #{nid}" in run["instruction"] and "5.3%" in run["instruction"] and "report_resolve_review_note" in run["instruction"]
    ev = [e for e in state.recent_events(store, 10) if e["kind"] == "steer"][-1]
    assert ev["label"] == "It was 5.3% in Q1 2026." and ev["detail"]["note"] == nid, "the feed shows their own words"
    with pytest.raises(LookupError):
        rt.reply_to_note(999999, "x")


# ---- the user is a source -----------------------------------------------------------------------------------------------


def _rs(store):
    rs = Store(url=store.url, schema=store.schema, report_id=f"fund-i-{Q3}")
    rs.ensure_report("Chui Ventures Fund I", "Q3 2026")
    return rs


def _fact(value, source, cell, label="Kenya GDP growth"):
    return Fact(label=label, value=value, unit="percent", source_file=source, source_cell=cell)


def test_a_figure_the_user_typed_is_accepted_when_the_quote_is_really_theirs(store):
    from chui_reporter.agent.ledger import verify_claim_ex

    rs = _rs(store)
    state.emit(store, "steer", "Kenya GDP growth was 5.3% in Q1 2026 according to the KNBS release.")
    ok, why, status = verify_claim_ex(_fact(5.3, "user:message", "Kenya GDP growth was 5.3% in Q1 2026"), None, rs)
    assert ok and status == "provided" and "stated by the user" in why
    ok, why, _ = verify_claim_ex(_fact(5.3, "user:message", "Kenya GDP growth was 9.9% in Q1 2026"), None, rs)
    assert not ok and "not in anything the user wrote" in why, "an invented quote is refused"
    ok, why, _ = verify_claim_ex(_fact(7.1, "user:message", "Kenya GDP growth was 5.3% in Q1 2026"), None, rs)
    assert not ok and "does not contain" in why, "the quote must contain the figure"
    ok, why, _ = verify_claim_ex(_fact(5.3, "user:message", "5.3"), None, rs)
    assert not ok and "quote the user's own words" in why


def test_the_users_answers_to_questions_count_as_their_word_too(store):
    from chui_reporter.agent.ledger import verify_claim_ex

    state.emit(store, "question.answered", "The policy rate is 8.75% as of June.")
    ok, _, status = verify_claim_ex(_fact(8.75, "user:message", "The policy rate is 8.75% as of June", "Kenya policy rate"), None, _rs(store))
    assert ok and status == "provided"


def test_a_figure_read_from_an_image_the_user_attached_is_accepted_as_theirs_and_flagged(tmp_path):
    from chui_reporter.agent.ledger import verify_claim_ex

    img = tmp_path / "Uploads" / "chart.png"
    img.parent.mkdir()
    img.write_bytes(b"\x89PNG")
    ok, why, status = verify_claim_ex(_fact(5.3, "chart.png", "bar for Kenya, top right"), lambda n: img, None)
    assert ok and status == "provided" and "flagged for your review" in why
    other = tmp_path / "elsewhere.png"
    other.write_bytes(b"\x89PNG")
    assert verify_claim_ex(_fact(5.3, "elsewhere.png", "x"), lambda n: other, None)[0] is False, "only images they attached"


def test_figures_in_the_files_they_attach_are_checked_in_those_files(tmp_path):
    from chui_reporter.agent.ledger import verify_claim_ex
    from docx import Document

    csv = tmp_path / "gdp.csv"
    csv.write_text("country,gdp\nKenya,5.3\n")
    ok, _, status = verify_claim_ex(_fact(5.3, "gdp.csv", ""), lambda n: csv, None)
    assert ok and status == "extracted"
    assert verify_claim_ex(_fact(6.6, "gdp.csv", ""), lambda n: csv, None)[0] is False
    doc = tmp_path / "note.docx"
    d = Document()
    d.add_paragraph("Kenya real GDP growth was 5.3 percent.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Nigeria", "3.89"
    d.save(str(doc))
    assert verify_claim_ex(_fact(5.3, "note.docx", ""), lambda n: doc, None)[0] is True
    assert verify_claim_ex(_fact(3.89, "note.docx", ""), lambda n: doc, None)[0] is True, "table cells are read too"
    assert verify_claim_ex(_fact(9.9, "note.docx", ""), lambda n: doc, None)[0] is False


def test_a_provided_figure_licenses_the_report_like_a_verified_one(store):
    rs = _rs(store)
    rs.add_facts([Fact(label="Kenya GDP growth", value=5.3, unit="percent", source_file="user:message", status="provided"),
                  Fact(label="Guess", value=7.7, unit="percent", source_file="x", status="claimed")])
    grounded = rs.grounded_values()
    assert 5.3 in grounded and 7.7 not in grounded


# ---- the Word file names no author -------------------------------------------------------------------------------------


def test_the_downloaded_word_file_carries_no_author():
    from docx import Document

    from chui_reporter.render.docmeta import blank_docx_author

    doc = Document()
    doc.add_paragraph("Hello")
    doc.core_properties.author = "python-docx"
    doc.core_properties.last_modified_by = "Somebody"
    doc.core_properties.comments = "generated"
    buf = io.BytesIO()
    doc.save(buf)
    cleaned = blank_docx_author(buf.getvalue())
    again = Document(io.BytesIO(cleaned))
    assert again.core_properties.author == "" and again.core_properties.last_modified_by == "" and again.core_properties.comments == ""
    assert [p.text for p in again.paragraphs] == ["Hello"], "the content is untouched"
    assert zipfile.ZipFile(io.BytesIO(cleaned)).testzip() is None
    assert blank_docx_author(b"not a zip") == b"not a zip"


def test_the_renderer_leaves_the_author_blank(store, tmp_path):
    from chui_reporter.render import report_writer as rw

    src = open(rw.__file__).read()
    assert 'core_properties.author = ""' in src and "Chui Ventures\"\n    doc.core_properties.author" not in src


# ---- the pages' calls ------------------------------------------------------------------------------------------------------


@pytest.fixture()
def client(store, rt):
    app = create_app(runtime=rt, auth=Auth("pw", "secret"))
    with TestClient(app, base_url="http://app.test") as c:
        assert c.post("/api/login", json={"password": "pw"}).status_code == 200
        yield c


def test_a_sync_for_a_quarter_the_server_has_left_is_refused_not_filed_under_the_wrong_one(client, rt):
    rt.set_period(Q3)
    data = b"hello"
    plan = {"manifest": [{"path": "Fund Financials/LP.xlsx", "sha256": _h(data), "size": 5}]}
    assert client.post("/api/sync/plan", json={**plan, "period": Q2}).status_code == 409
    assert client.post("/api/sync/plan", json={**plan, "period": Q3}).json()["need"] == ["Fund Financials/LP.xlsx"]
    assert client.put(f"/api/sync/file?path=Fund%20Financials/LP.xlsx&sha256={_h(data)}&period={Q2}", content=data).status_code == 409
    assert client.put(f"/api/sync/file?path=Fund%20Financials/LP.xlsx&sha256={_h(data)}&period={Q3}", content=data).json() == {"stored": True}
    r = client.post("/api/sync/commit", json={**plan, "period": Q3, "folder": "q3 data"})
    assert r.status_code == 200 and r.json()["changes"]["added"] == ["Fund Financials/LP.xlsx"]
    s = client.get("/api/state").json()
    assert s["workspace"]["files"] == 1 and s["workspace"]["folder"] == "q3 data" and s["workspace"]["last_sync_at"]


def test_clearing_and_moving_a_quarters_files_from_the_page(client, rt, store):
    rt.set_period(Q3)
    data = b"hello"
    client.put(f"/api/sync/file?path=Fund%20Financials/LP.xlsx&sha256={_h(data)}&period={Q3}", content=data)
    client.post("/api/sync/commit", json={"manifest": [{"path": "Fund Financials/LP.xlsx", "sha256": _h(data), "size": 5}],
                                          "period": Q3, "folder": "oops"})
    moved = client.post("/api/sources/move", json={"to": "2026Q2"}).json()
    assert moved["moved"] == 1 and moved["state"]["workspace"]["files"] == 0
    assert client.post("/api/sources/move", json={"to": "2026Q3"}).status_code == 400, "not onto itself"
    assert client.post("/api/sources/move", json={"to": "garbage"}).status_code == 400
    assert _paths(store, Q2) == ["Fund Financials/LP.xlsx"]
    rt.set_period(Q2)
    cleared = client.post("/api/sources/clear", json={}).json()
    assert cleared["removed"] == 1 and cleared["state"]["status"]["phase"] == "empty"
    rid = rt._create("build", "x")
    state.update_run(store, rid, status="running")
    assert client.post("/api/sources/clear", json={}).status_code == 409
    assert client.post("/api/period", json={"period": "2026Q3"}).status_code == 409


def test_the_period_endpoint_gives_the_new_quarters_state(client):
    s = client.post("/api/period", json={"period": "Q3 2026"}).json()
    assert s["period"]["code"] == Q3 and s["version"] is None
    assert client.post("/api/period", json={"period": "nonsense"}).status_code == 400


def test_review_notes_are_resolved_and_reopened_from_the_page(client, rt):
    rs = rt.report_store()
    rs.add_review_note("Valuation", "Two marks are stale.", "warning")
    notes = client.get("/api/notes").json()["notes"]
    nid = notes[0]["id"]
    assert notes[0]["status"] == "open"
    r = client.post(f"/api/notes/{nid}/resolve", json={"resolution": "Fine as it is"}).json()["note"]
    assert r["status"] == "resolved" and r["resolution"] == "Fine as it is" and r["resolved_at"]
    assert client.post(f"/api/notes/{nid}/reopen").json()["note"]["status"] == "open"
    assert client.post("/api/notes/424242/resolve", json={}).status_code == 404
    reply = client.post("/api/steer", json={"text": "They are fine, I checked.", "note": nid}).json()
    assert reply["ok"] and client.get("/api/notes").json()["notes"][0]["status"] == "answered"
    assert client.post("/api/steer", json={"text": "x", "note": 424242}).status_code == 404


def test_the_word_download_has_no_author_even_if_it_was_stored_with_one(client, rt):
    from docx import Document

    rs = rt.report_store()
    doc = Document()
    doc.core_properties.author = "Chui Ventures"
    buf = io.BytesIO()
    doc.save(buf)
    state.save_version(rs, 1, {"sections": []}, b"%PDF-1", buf.getvalue(), None)
    r = client.get("/api/report/download?fmt=docx")
    assert r.status_code == 200
    assert Document(io.BytesIO(r.content)).core_properties.author == ""


# ---- a conversation is not work ------------------------------------------------------------------------------------------


def _conversation_runtime(store, script):
    from langchain_core.messages import AIMessage  # noqa: F401  (scripts use it)

    from tests.test_runtime import _runtime

    return _runtime(store, script)


def test_a_greeting_gets_an_answer_and_is_not_sent_off_to_render_a_report(store):
    from langchain_core.messages import AIMessage

    rt, llm = _conversation_runtime(store, [AIMessage(content="Hello! Tell me what you need.")])
    rt.set_period(Q3)
    rid = rt._create("steer", "hey")
    assert rt.execute(rid) == "done"
    assert len(llm.seen) == 1, "one model call: the answer, nothing more"
    kinds = [e["kind"] for e in state.recent_events(store, 100)]
    assert "nudge" not in kinds and "note" in kinds
    note = [e for e in state.recent_events(store, 100) if e["kind"] == "note"][-1]
    assert note["label"] == "Hello! Tell me what you need."


def test_a_message_that_changes_the_report_still_has_to_be_rendered_and_checked(store):
    from tests.fakes import call
    from tests.test_runtime import _events

    rt, _ = _conversation_runtime(store, [call("write_it", {"key": "1.1", "text": "A new paragraph."}, "a")])
    rt.set_period(Q3)
    rid = rt._create("steer", "change 1.1")
    rt.execute(rid)
    assert _events(store, "nudge"), "work on the report is sent on to render and inspect it"


def test_a_build_with_nothing_written_is_still_sent_back_to_work(store):
    from langchain_core.messages import AIMessage
    from tests.test_runtime import _events

    rt, _ = _conversation_runtime(store, [AIMessage(content="I am done.")])
    rt.set_period(Q3)
    rt.execute(rt._create("build", "Build the report"))
    assert _events(store, "nudge"), "a build run is held to the render-and-inspect check"


# ---- what the last sync brought in ------------------------------------------------------------------------------------------


def test_each_sync_that_changes_something_leaves_a_summary_the_page_can_show(store, rt):
    rt.set_period(Q3)
    m = [_put(store, Q3, "Fund Financials/LP.xlsx", b"lp"), _put(store, Q3, "Valuation Reports/A.xlsx", b"a")] + _brand(store, Q3)
    ch = sync.commit(store, m, Q3, folder="q3 data", skipped=["Valuation Reports/huge.xlsx"])
    ls = rt.snapshot_state()["workspace"]["last_sync"]
    assert ls["folder"] == "q3 data" and ls["files"] == 6 and ls["added"] == len(ch.added) == 6 and ls["brand"] == 4
    assert ls["modified"] == 0 and ls["removed"] == 0 and ls["skipped"] == ["Valuation Reports/huge.xlsx"] and ls["at"]
    sync.commit(store, m, Q3, folder="q3 data")                       # the watcher looking again: nothing changed
    assert rt.snapshot_state()["workspace"]["last_sync"] == ls, "an unchanged look does not wipe the receipt"
    m[0] = _put(store, Q3, "Fund Financials/LP.xlsx", b"lp, revised")
    sync.commit(store, m, Q3, folder="q3 data")
    again = rt.snapshot_state()["workspace"]["last_sync"]
    assert again["modified"] == 1 and again["added"] == 0 and again["at"] >= ls["at"]


def test_the_sync_is_announced_in_the_feed_with_the_folder_and_the_counts(store, rt):
    rt.set_period(Q3)
    ch = sync.Changes(added=["a", "b", "c"], modified=["d"], removed=["e"])
    assert rt._sync_label(ch, "q3 data") == "Synced 4 files from “q3 data”: 3 new, 1 updated, 1 removed"
    assert rt._sync_label(sync.Changes(removed=["x"])) == "Folder changed: 1 removed"
    rt.on_sync(sync.Changes(added=["Fund Financials/LP.xlsx"]), "q3 data")
    ev = [e for e in state.recent_events(store, 20) if e["kind"] == "source.sync"][-1]
    assert ev["label"] == "Synced 1 file from “q3 data”: 1 new" and ev["detail"]["folder"] == "q3 data"


def test_the_commit_call_answers_with_what_the_page_needs_to_say_sync_complete(client, rt):
    rt.set_period(Q3)
    data = b"hello"
    client.put(f"/api/sync/file?path=Fund%20Financials/LP.xlsx&sha256={_h(data)}&period={Q3}", content=data)
    r = client.post("/api/sync/commit", json={"manifest": [{"path": "Fund Financials/LP.xlsx", "sha256": _h(data), "size": 5}],
                                              "period": Q3, "folder": "q3 data", "skipped": ["big.xlsx"]}).json()
    assert r["folder"] == "q3 data" and r["files"] == 1 and r["brand"] == 0 and r["changes"]["added"] == ["Fund Financials/LP.xlsx"]
    assert client.get("/api/state").json()["workspace"]["last_sync"]["skipped"] == ["big.xlsx"]


# ---- the operator's reset ----------------------------------------------------------------------------------------------------


def _populate(store, rt, code, report_note="n"):
    from chui_reporter import admin  # noqa: F401

    rt.set_period(code)
    rs = rt.report_store()
    rs.set_section("1.1", "Overview", "text", 1)
    rs.set_table("t1", "A table", ["a"], [["1"]], "1.1", {})
    rs.add_facts([Fact(label=f"{code} fact", value=1.0, unit="USD", source_file="f.xlsx", source_cell="A1")])
    rs.add_review_note("Area", f"{report_note} {code}", "info")
    state.save_version(rs, 3, {"sections": []}, b"%PDF", None, None)
    sync.commit(store, [_put(store, code, "Fund Financials/LP.xlsx", code.encode())], code, folder=f"{code} folder")
    rid = rt._create("build", "x")
    state.update_run(store, rid, status="done")
    state.emit(store, "step", f"{code} step")
    return rs


def test_resetting_a_quarter_returns_it_to_blank_and_touches_no_other_quarter(store, rt):
    from chui_reporter import admin

    for code in (Q2, Q3):
        _populate(store, rt, code)
    sync.commit(store, [_put(store, Q3, "Fund Financials/LP.xlsx", Q3.encode())] + _brand(store, Q3), Q3)      # Q3's folder, brand kit included
    store_q2 = Store(url=store.url, schema=store.schema, report_id=f"fund-i-{Q2}")
    from chui_reporter.services import web

    web.save_snapshot(store, "https://cbk.example/x", "page text " * 30, "p", "direct")

    dry = admin.reset_quarter(store, Q2)
    assert dry["applied"] is False and dry["counts"]["files"] >= 1 and dry["counts"]["runs"] == 1
    assert state.latest_version(store_q2) is not None, "a dry run deletes nothing"

    done = admin.reset_quarter(store, Q2, apply=True)
    assert done["applied"] is True and done["counts"]["runs"] == 1 and done["counts"]["sessions"] >= 1
    assert state.latest_version(store_q2) is None and store_q2.sections() == [] and store_q2.review_notes() == []
    assert _paths(store, Q2) == [] and state.workspace(store, Q2)["last_sync_at"] is None
    assert not [e for e in state.recent_events(store, 200) if e["label"] == f"{Q2} step"]
    # the other quarter and the shared brand kit are exactly as they were
    store_q3 = Store(url=store.url, schema=store.schema, report_id=f"fund-i-{Q3}")
    assert state.latest_version(store_q3) is not None and len(store_q3.sections()) == 1 and len(store_q3.review_notes()) == 1
    assert _paths(store, Q3) == ["Fund Financials/LP.xlsx"] and len(_paths(store, sync.SHARED)) == 4
    assert [e for e in state.recent_events(store, 200) if e["label"] == f"{Q3} step"]
    with store.conn() as c:
        assert c.execute(f"SELECT 1 FROM {store._t('web_snapshot')}").fetchone() is None, "the web cache was emptied"
    # and the quarter works again from blank
    rt.set_period(Q2)
    assert rt.snapshot_state()["status"]["phase"] == "empty" and rt.snapshot_state()["version"] is None


def test_the_brand_kit_goes_only_when_asked_and_a_reset_refuses_while_the_agent_is_working(store, rt):
    from chui_reporter import admin

    _populate(store, rt, Q2)
    sync.commit(store, _brand(store, Q2), Q2)
    admin.reset_quarter(store, Q2, apply=True)
    assert len(_paths(store, sync.SHARED)) == 4
    _populate(store, rt, Q2)
    admin.reset_quarter(store, Q2, brand=True, apply=True)
    assert _paths(store, sync.SHARED) == []
    rid = rt._create("build", "x")
    state.update_run(store, rid, status="running")
    with pytest.raises(admin.Busy):
        admin.reset_quarter(store, Q2, apply=True)


def test_a_reset_that_fails_part_way_changes_nothing(store, rt, monkeypatch):
    from chui_reporter import admin

    _populate(store, rt, Q2)
    store_q2 = Store(url=store.url, schema=store.schema, report_id=f"fund-i-{Q2}")
    real = admin.pr.Period.parse

    def parse(text):
        return real(text)

    # make a late delete fail: the table the reset empties last before the dry-run check
    monkeypatch.setattr(admin, "REPORT_TABLES", ("section", "no_such_table"))
    with pytest.raises(Exception):
        admin.reset_quarter(store, Q2, apply=True)
    assert len(store_q2.sections()) == 1 and state.latest_version(store_q2) is not None, "rolled back as a whole"


# ---- the feed speaks in names a person would use ----------------------------------------------------------------------------


def test_the_feed_names_files_the_way_a_person_would():
    from chui_reporter.runtime.narrator import describe_call, friendly_file

    names = ["06. LAMI-Portfolio Valuation Report .xlsx", "12. OneHealth-Portfolio Valuation Report .xlsx",
             "Fund Model Cap $16.3 M (Q2 2026) (2).xlsx", "06_30_2026 - CHUI VENTURES FUND I, LP Financial Package (1).pdf"]
    assert friendly_file(names[0]) == "LAMI valuation report" and friendly_file(names[1]) == "OneHealth valuation report"
    assert friendly_file(names[2]) == "Fund Model Cap $16.3 M (Q2 2026)"
    assert describe_call("excel_sheets", {"file_name": "06"}, names)[1] == "Opening LAMI valuation report"
    assert describe_call("excel_sheets", {"file_name": "Fund Model"}, names)[1] == "Opening Fund Model Cap $16.3 M (Q2 2026)"
    assert describe_call("read_pdf", {"file_name": "Financial Package"}, names)[1].startswith("Reading 06_30_2026 - CHUI VENTURES FUND I")
    assert describe_call("excel_sheets", {"file_name": "Uncover"}, [])[1] == "Opening Uncover", "without the names, as before"
